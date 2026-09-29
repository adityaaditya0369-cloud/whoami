"""Validation records, the go/no-go gate, cutover and rollback.

State path (every step audited, named people where it matters):
  APPROVED -> PF_CONFIGURED        created on PingFederate (gated write or manual import, confirmed by sync)
  PF_CONFIGURED -> TESTING         first pre-cutover validation recorded
  TESTING -> BUSINESS_VALIDATION   app owner signs off (needs a passing validation)
  BUSINESS_VALIDATION -> READY_FOR_CUTOVER   go decision (gate must be green)
  READY_FOR_CUTOVER -> MIGRATED    SP switched to PingFederate
  MIGRATED -> VALIDATED            post-cutover validation passes (automatic)
  MIGRATED -> FAILED               post-cutover validation fails (automatic)
  any test/cutover state -> FAILED declared by a person (VALIDATED too, within the rollback window)
  FAILED -> ROLLBACK -> OKTA_ACTIVE   rollback started, then SP confirmed back on Okta
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import timedelta
from urllib.parse import quote

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.models.db import Application, MigrationPlan, ValidationResult, utcnow
from app.models.state import MigrationState
from app.services import saml_validation as sv
from app.services.state_service import record_event, transition

S = MigrationState
DEFAULT_CTX = "urn:oasis:names:tc:SAML:2.0:ac:classes:PasswordProtectedTransport"
STAGES = {"BASELINE_OKTA": "Okta baseline", "PRE_CUTOVER": "Pre-cutover test", "POST_CUTOVER": "Post-cutover check"}
ROLLBACK_WINDOW_DAYS = 7
FAILABLE = {S.PF_CONFIGURED, S.VENDOR_CONFIGURED, S.TESTING, S.BUSINESS_VALIDATION, S.READY_FOR_CUTOVER, S.MIGRATED}


def _named(actor: str) -> str:
    actor = (actor or "").strip()
    if not actor or actor.lower() in {"system", "ai", "claude"}:
        raise ValueError("A named person is required")
    return actor


def _acs(app: Application) -> list[str]:
    s = app.saml
    urls = [e.get("url") for e in (s.acs_endpoints or []) if e.get("url")]
    return urls or ([s.sso_acs_url] if s.sso_acs_url else [])


def okta_issuer(app: Application) -> str | None:
    s = app.saml
    return s.idp_issuer if s.idp_issuer and "${" not in s.idp_issuer else s.metadata_entity_id


def _okta_cert(app: Application):
    certs = [c for c in app.certificates if c.is_active_signing_key] or list(app.certificates)
    return certs[0] if certs else None


def _common(app: Application) -> dict:
    s = app.saml
    codes = {f.code for f in app.findings}
    return {"audience": s.audience, "acs_urls": _acs(app), "name_id_format": s.name_id_format,
            "attributes": [c.name for c in app.claims],
            "group_attributes": [c.name for c in app.claims if c.claim_type == "GROUP"],
            "authn_context": s.authn_context_class_ref if "NON_DEFAULT_AUTHN_CONTEXT" in codes else None}


def expected_values(app: Application, settings: Settings) -> dict:
    """What a PingFederate response for this app must look like."""
    s = app.saml
    codes = {f.code for f in app.findings}
    return {**_common(app),
            "issuer": s.idp_issuer if "CUSTOM_IDP_ISSUER" in codes else settings.pf_entity_id,
            "sign_assertion": bool(s.assertion_signed), "signature_algorithm": "rsa-sha256"}


def okta_expected_values(app: Application) -> dict:
    s = app.saml
    c = _okta_cert(app)
    alg = "rsa-sha1" if (s.signature_algorithm or "").upper().endswith("SHA1") else "rsa-sha256"
    return {**_common(app), "issuer": okta_issuer(app), "sign_assertion": bool(s.assertion_signed),
            "signature_algorithm": alg, "cert_pem": sv.pem_of(c.x5c) if c and c.x5c else None}


def rollback_values(app: Application, settings: Settings) -> dict:
    """What the SP must be set back to if the cutover fails (captured from Okta)."""
    s, c = app.saml, _okta_cert(app)
    org = settings.okta_org
    return {"idp_entity_id": okta_issuer(app), "sso_url": s.metadata_sso_url,
            "slo_url": s.slo_logout_url and (f"{org}/app/{app.okta_name}/{app.id}/slo/saml" if org else None),
            "metadata_url": f"{org}/app/{app.id}/sso/saml/metadata" if org else None,
            "certificate_sha256": c.sha256_thumbprint if c else None, "certificate_sha1": c.sha1_thumbprint if c else None,
            "certificate_not_after": c.not_after.strftime("%Y-%m-%d") if c and c.not_after else None,
            "certificate_pem": sv.pem_of(c.x5c) if c and c.x5c else None,
            "okta_app_status": app.okta_status}


def cutover_values(app: Application, settings: Settings) -> dict:
    s = app.saml
    base = settings.pf_base_url.rstrip("/")
    exp = expected_values(app, settings)
    ent = quote(s.audience or "", safe="")
    return {"idp_entity_id": exp["issuer"], "sso_url": f"{base}/idp/SSO.saml2",
            "slo_url": f"{base}/idp/SLO.saml2" if s.slo_enabled else None,
            "metadata_url": f"{base}/pf/federation_metadata.ping?PartnerSpId={ent}" if s.audience else None,
            "idp_initiated_url": f"{base}/idp/startSSO.ping?PartnerSpId={ent}" if s.audience else None,
            "certificate": "PingFederate signing certificate (from the metadata)"}


# --- validation records -------------------------------------------------------------
def results(session: Session, app_id: str) -> list[ValidationResult]:
    return list(session.scalars(select(ValidationResult).where(ValidationResult.app_id == app_id)
                                .order_by(ValidationResult.id.desc())))


def latest(session: Session, app_id: str, stage: str) -> ValidationResult | None:
    return session.scalars(select(ValidationResult).where(ValidationResult.app_id == app_id,
                                                          ValidationResult.stage == stage)
                           .order_by(ValidationResult.id.desc()).limit(1)).first()


def default_stage(app: Application) -> str:
    if app.state in (S.MIGRATED.value, S.VALIDATED.value):
        return "POST_CUTOVER"
    if app.state in (S.PF_CONFIGURED.value, S.VENDOR_CONFIGURED.value, S.TESTING.value,
                     S.BUSINESS_VALIDATION.value, S.READY_FOR_CUTOVER.value, S.FAILED.value):
        return "PRE_CUTOVER"
    return "BASELINE_OKTA"


def _last_transition(session: Session, app_id: str, to_state: S):
    from app.models.db import MigrationEvent
    return session.scalars(select(MigrationEvent).where(MigrationEvent.app_id == app_id,
                                                        MigrationEvent.event_type == "STATE_CHANGE",
                                                        MigrationEvent.to_state == to_state.value)
                           .order_by(MigrationEvent.id.desc()).limit(1)).first()


def _signing_cert(settings: Settings, stage: str, app: Application, pasted: str | None) -> tuple[str | None, str | None]:
    """(certificate PEM, where it came from). The certificate inside the response is never trusted.

    Gating stages use the configured PingFederate certificate. A pasted certificate is only used when none
    is configured, and such a result never opens the go/no-go gate (cert_source == "pasted")."""
    if stage == "BASELINE_OKTA":  # a comparison reference only; it never opens the gate
        if pasted:
            return pasted, "pasted"
        pem = okta_expected_values(app).get("cert_pem")
        return (pem, "okta_discovery") if pem else (None, None)
    f = settings.pf_signing_cert_file
    if f:
        if not f.exists():
            raise ValueError(f"PF_SIGNING_CERT_FILE is set but not found: {f}")
        return f.read_text(), "configured"
    return (pasted, "pasted") if pasted else (None, None)


def issued_at(vr: ValidationResult):
    return sv._dt((vr.result or {}).get("summary", {}).get("issue_instant")) or vr.created_at


def record_validation(session: Session, app: Application, settings: Settings, response: str, actor: str,
                      stage: str, source: str = "CAPTURED", cert_pem: str | None = None) -> ValidationResult:
    """Validate a response and store the result (not the raw response). May move the state."""
    actor = (actor or "").strip()
    if not actor:
        raise ValueError("Your name is required")
    if stage not in STAGES:
        raise ValueError("Unknown stage")
    if not app.is_saml or app.saml is None:
        raise ValueError("SAML validation applies to SAML apps")
    pem, cert_source = _signing_cert(settings, stage, app, cert_pem)
    gating = stage != "BASELINE_OKTA"
    if stage == "BASELINE_OKTA":
        expected, baseline = okta_expected_values(app), None
    else:
        expected = expected_values(app, settings)
        b = latest(session, app.id, "BASELINE_OKTA")
        baseline = b.result.get("summary") if b and b.verdict != "ERROR" else None
    rep = sv.validate(response, expected, baseline=baseline, signing_cert_pem=pem,
                      captured=(source == "CAPTURED"), skew_seconds=settings.validation_clock_skew_seconds,
                      strict=gating)
    # The same response cannot be recorded twice (e.g. a pre-cutover response re-used as post-cutover proof).
    dup = session.scalars(select(ValidationResult).where(ValidationResult.app_id == app.id,
                                                         ValidationResult.response_sha256 == rep.response_sha256)
                          .limit(1)).first()
    if dup:
        raise ValueError(f"This exact response was already recorded (validation #{dup.id}). Capture a new sign-in.")
    issued = sv._dt(rep.summary.get("issue_instant"))
    if stage == "POST_CUTOVER":
        mig = _last_transition(session, app.id, S.MIGRATED)
        if mig and (issued is None or issued < mig.ts - timedelta(seconds=settings.validation_clock_skew_seconds)):
            raise ValueError("This response was issued before the cutover was recorded; capture a sign-in made after "
                             f"the cutover ({mig.ts:%d %b %Y %H:%M} UTC)")
    if cert_source == "pasted":
        rep.checks.append(sv.Check("cert_source", "Certificate source", "configured (PF_SIGNING_CERT_FILE)",
                                   "pasted in the form", sv.WARN,
                                   "Useful for a quick check, but only a configured certificate counts for the gate"))
    from app.services.approval import approved_plan
    plan = approved_plan(session, app.id)
    vr = ValidationResult(app_id=app.id, result={**rep.as_dict(), "expected": {k: v for k, v in expected.items()
                                                                                 if k != "cert_pem"},
                                                  "baseline_used": bool(baseline), "cert_source": cert_source,
                                                  "cert_sha256": sv.cert_sha256(sv.pem_of(pem)) if pem else None},
                          verdict=rep.verdict, stage=stage, source=source, actor=actor,
                          response_sha256=rep.response_sha256, plan_id=plan.id if plan else None)
    session.add(vr)
    session.flush()
    record_event(session, "SAML_VALIDATION", actor, app.id,
                 {"validation_id": vr.id, "stage": stage, "source": source, "verdict": vr.verdict,
                  "cert_source": cert_source, "failed": [c.key for c in rep.checks if c.status == sv.FAIL]})
    _after_validation(session, app, vr, settings)
    return vr


def counts_for_gate(vr: ValidationResult | None) -> bool:
    """A passing result verified with the configured PingFederate certificate."""
    return bool(vr and vr.verdict in ("PASS", "PASS_WITH_WARNINGS")
                and (vr.result or {}).get("cert_source") == "configured")


def _after_validation(session: Session, app: Application, vr: ValidationResult, settings: Settings) -> None:
    st = S(app.state)
    ok = counts_for_gate(vr)
    if vr.stage == "PRE_CUTOVER" and st in (S.PF_CONFIGURED, S.VENDOR_CONFIGURED):
        transition(session, app, S.TESTING, "system", reason=f"Pre-cutover validation #{vr.id} recorded ({vr.verdict})")
    elif vr.stage == "POST_CUTOVER" and st == S.MIGRATED:
        if ok:
            transition(session, app, S.VALIDATED, "system", reason=f"Post-cutover validation #{vr.id} passed")
        elif vr.verdict == "FAIL":
            transition(session, app, S.FAILED, "system",
                       reason=f"Post-cutover validation #{vr.id} failed - decide on rollback",
                       detail={"validation_id": vr.id})


def fresh(vr: ValidationResult | None, settings: Settings) -> bool:
    """Recent = the response itself was issued recently (not just recorded recently)."""
    return bool(vr and issued_at(vr) >= utcnow() - timedelta(days=settings.validation_max_age_days))


def current_pre(session: Session, app: Application) -> ValidationResult | None:
    """The pre-cutover validation that counts: newer than the last failure and made against the current plan."""
    from app.services.approval import approved_plan
    plan = approved_plan(session, app.id)
    failed = _last_transition(session, app.id, S.FAILED)
    for vr in results(session, app.id):
        if vr.stage != "PRE_CUTOVER":
            continue
        if failed and vr.created_at <= failed.ts:
            return None
        if plan and vr.plan_id != plan.id:
            return None
        return vr
    return None


# --- gate ------------------------------------------------------------------------
@dataclass
class GateCheck:
    label: str
    ok: bool
    detail: str
    required: bool = True


def gate(session: Session, app: Application, settings: Settings) -> list[GateCheck]:
    from app.services.approval import approved_plan
    from app.services.reconcile import reconcile
    plan = approved_plan(session, app.id)
    rec = next((r for r in reconcile(session)[0] if r.app.id == app.id), None)
    pre = current_pre(session, app)
    base = latest(session, app.id, "BASELINE_OKTA")
    rb = rollback_values(app, settings)
    checks = [
        GateCheck("Plan approved", bool(plan), f"v{plan.version} by {plan.decided_by}" if plan else "No approved plan"),
        GateCheck("Connection found on PingFederate", bool(rec and rec.pf),
                  (f"{rec.pf.name} ({'active' if rec.pf.active else 'disabled - enable it at cutover'})"
                   if rec and rec.pf else "Not found - run a PingFederate sync")),
        GateCheck("No differences to the plan on PingFederate", bool(rec and rec.pf and not rec.problems),
                  ", ".join(c.check for c in rec.problems) if rec and rec.problems else ("Matches" if rec and rec.pf else "—")),
        GateCheck("Pre-cutover validation passed (current plan, since any failure)", counts_for_gate(pre),
                  (f"#{pre.id} {pre.verdict} on {pre.created_at:%d %b %H:%M}"
                   + ("" if (pre.result or {}).get("cert_source") == "configured"
                      else " - signature not verified with the configured PingFederate certificate"))
                  if pre else "None recorded (a new test is needed after a failure or a new plan)"),
        GateCheck(f"Tested sign-in is recent (≤ {settings.validation_max_age_days} days)", fresh(pre, settings),
                  f"issued {issued_at(pre):%d %b %Y}" if pre else "—"),
        GateCheck("Business sign-off", app.state in (S.BUSINESS_VALIDATION.value, S.READY_FOR_CUTOVER.value),
                  "Signed off" if app.state in (S.BUSINESS_VALIDATION.value, S.READY_FOR_CUTOVER.value) else "Not yet"),
        GateCheck("Rollback values captured", bool(rb["idp_entity_id"] and rb["certificate_sha256"]),
                  "Okta issuer and signing certificate on file" if rb["idp_entity_id"] and rb["certificate_sha256"]
                  else "Missing Okta issuer or certificate"),
        GateCheck("Okta app still active (rollback possible)", app.okta_status == "ACTIVE", app.okta_status),
        GateCheck("Okta baseline captured for comparison", bool(base), f"#{base.id} {base.verdict}" if base else
                  "Recommended: capture an Okta response for the test user", required=False),
    ]
    if rb["certificate_not_after"]:
        from datetime import datetime
        exp = datetime.strptime(rb["certificate_not_after"], "%Y-%m-%d")
        checks.append(GateCheck("Okta certificate valid through the rollback window", exp > utcnow() + timedelta(days=7),
                                f"expires {rb['certificate_not_after']}"))
    return checks


def gate_ok(checks: list[GateCheck]) -> bool:
    return all(c.ok for c in checks if c.required)


# --- actions ---------------------------------------------------------------------
def business_signoff(session: Session, app: Application, settings: Settings, actor: str, comment: str) -> None:
    actor = _named(actor)
    pre = current_pre(session, app)
    if not counts_for_gate(pre):
        raise ValueError("A passing pre-cutover validation (current plan, verified with the configured PingFederate "
                         "certificate, recorded after any failure) is needed before business sign-off")
    transition(session, app, S.BUSINESS_VALIDATION, actor, reason=comment or "Business validation signed off",
               detail={"validation_id": pre.id})


def go_decision(session: Session, app: Application, settings: Settings, actor: str, comment: str) -> None:
    actor = _named(actor)
    checks = gate(session, app, settings)
    if not gate_ok(checks):
        raise ValueError("Go/no-go gate is not green: " + "; ".join(c.label for c in checks if c.required and not c.ok))
    transition(session, app, S.READY_FOR_CUTOVER, actor, reason=comment or "Go decision",
               detail={"gate": [c.label for c in checks if c.ok]})


def cutover_done(session: Session, app: Application, actor: str, comment: str, settings: Settings | None = None) -> None:
    actor = _named(actor)
    if not comment.strip():
        raise ValueError("Record what was switched and when (e.g. change number)")
    if settings is not None:
        bad = [c.label for c in gate(session, app, settings) if c.required and not c.ok
               and c.label != "Business sign-off"]
        if bad:
            raise ValueError("The go/no-go gate changed since the go decision: " + "; ".join(bad))
    transition(session, app, S.MIGRATED, actor, reason=comment)


def declare_failed(session: Session, app: Application, actor: str, reason: str) -> None:
    actor = _named(actor)
    if not reason.strip():
        raise ValueError("A reason is required")
    if S(app.state) == S.VALIDATED:
        mig = _last_transition(session, app.id, S.MIGRATED)
        if not mig or mig.ts < utcnow() - timedelta(days=ROLLBACK_WINDOW_DAYS):
            raise ValueError(f"The {ROLLBACK_WINDOW_DAYS}-day rollback window after cutover has closed")
    elif S(app.state) not in FAILABLE:
        raise ValueError(f"Cannot declare failure from {app.state}")
    transition(session, app, S.FAILED, actor, reason=reason)


def retry_testing(session: Session, app: Application, actor: str, reason: str) -> None:
    actor = _named(actor)
    transition(session, app, S.TESTING, actor, reason=reason or "Fixed; testing again")


def start_rollback(session: Session, app: Application, settings: Settings, actor: str, reason: str) -> None:
    actor = _named(actor)
    if not reason.strip():
        raise ValueError("A reason is required")
    transition(session, app, S.ROLLBACK, actor, reason=reason,
               detail={"restore": {k: v for k, v in rollback_values(app, settings).items() if k != "certificate_pem"}})


def confirm_rollback(session: Session, app: Application, actor: str, comment: str) -> None:
    actor = _named(actor)
    if not comment.strip():
        raise ValueError("Confirm how you checked that users sign in through Okta again")
    transition(session, app, S.OKTA_ACTIVE, actor, reason=comment)


# --- runbook -----------------------------------------------------------------------
def runbook(session: Session, app: Application, settings: Settings) -> dict:
    from app.services.approval import approved_plan
    plan = approved_plan(session, app.id)
    cv, rb = cutover_values(app, settings), rollback_values(app, settings)
    owner = app.business_owner or "App owner (not captured)"
    date = (plan.plan.get("cutover_date") if plan else None) or "to be agreed"
    return {
        "app": app.label, "owner": owner, "cutover_date": date, "plan": f"v{plan.version}" if plan else "not approved",
        "pre_checks": [f"{c.label}: {c.detail}" for c in gate(session, app, settings)],
        "cutover_steps": [
            f"Freeze: no Okta or PingFederate changes for {app.label} from T-24h.",
            "Enable the SP connection on PingFederate (it was created disabled).",
            f"App owner / vendor sets the IdP on the SP: entity ID {cv['idp_entity_id']}, SSO URL {cv['sso_url']}"
            + (f", SLO URL {cv['slo_url']}" if cv["slo_url"] else "") + ", PingFederate signing certificate"
            + (f" (metadata: {cv['metadata_url']})" if cv["metadata_url"] else "") + ".",
            "A test user signs in (SP-initiated and IdP-initiated); capture the SAML response.",
            "Record it as a post-cutover validation; the tool compares it with the plan and the Okta baseline.",
            "Pilot users confirm sign-in, roles and group access.",
            "Update portal links and bookmarks to the PingFederate IdP-initiated URL"
            + (f" ({cv['idp_initiated_url']})" if cv["idp_initiated_url"] else "") + ".",
        ],
        "go_no_go": ["Post-cutover validation PASS (or PASS WITH WARNINGS accepted by the IAM lead)",
                     "No sign-in errors reported by pilot users within 30 minutes",
                     "App owner confirms roles/permissions are correct"],
        "rollback_triggers": ["Post-cutover validation FAIL (NameID, audience or signature)",
                              "Users cannot sign in or land on the wrong account",
                              "Roles or group access wrong and no fix within the change window"],
        "rollback_steps": [
            "Declare the failure in the tool (reason recorded), then start the rollback.",
            f"App owner / vendor sets the IdP on the SP back to Okta: entity ID {rb['idp_entity_id']}, "
            f"SSO URL {rb['sso_url'] or '(from Okta metadata)'}"
            + (f", metadata {rb['metadata_url']}" if rb["metadata_url"] else "") + ".",
            f"Okta signing certificate: SHA-256 {rb['certificate_sha256'] or '—'} (valid until {rb['certificate_not_after'] or '—'}).",
            "Disable the SP connection on PingFederate (do not delete it - keep it for the retry).",
            "A test user signs in through Okta; confirm in the tool (moves the app to Okta active).",
            "Record what failed; fix, re-test and plan a new cutover.",
        ],
        "keep": "Keep the Okta app active until the rollback window (7 days after cutover) has closed.",
        "cutover_values": cv, "rollback_values": {k: v for k, v in rb.items() if k != "certificate_pem"},
    }


def runbook_docx(rb: dict, customer: str) -> bytes:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    d.add_paragraph(f"Cutover and rollback runbook: {rb['app']}", style="Title")
    d.add_paragraph(f"{customer} · owner {rb['owner']} · cutover {rb['cutover_date']} · plan {rb['plan']}")
    for title, key, style in (("Pre-checks (go/no-go gate)", "pre_checks", "List Bullet"),
                              ("Cutover steps", "cutover_steps", "List Number"),
                              ("Go / no-go after cutover", "go_no_go", "List Bullet"),
                              ("Rollback triggers", "rollback_triggers", "List Bullet"),
                              ("Rollback steps", "rollback_steps", "List Number")):
        d.add_heading(title, level=1)
        for line in rb[key]:
            d.add_paragraph(line, style=style)
    d.add_heading("Values", level=1)
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Setting", "Cutover (PingFederate)", "Rollback (Okta)"
    for label, a, b in (("IdP entity ID", "idp_entity_id", "idp_entity_id"), ("SSO URL", "sso_url", "sso_url"),
                        ("SLO URL", "slo_url", "slo_url"), ("Metadata", "metadata_url", "metadata_url"),
                        ("Signing certificate", "certificate", "certificate_sha256")):
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text = label, str(rb["cutover_values"].get(a) or "—"), str(rb["rollback_values"].get(b) or "—")
    d.add_paragraph(rb["keep"])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def plan_of(session: Session, vr: ValidationResult) -> MigrationPlan | None:
    return session.get(MigrationPlan, vr.plan_id) if vr.plan_id else None
