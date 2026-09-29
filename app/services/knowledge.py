"""Knowledge base for the review agents (the "RAG agent").

People upload the customer's runbooks, vendor SSO guides and migration notes.
Text is split into chunks and searched with BM25 - a keyword ranking that runs
locally, so documents never leave the machine for indexing (Anthropic has no
embeddings API and we do not want a third-party one for customer documents).
Agents get the top passages with a reference (KB-<id>) and must cite only
those references; the policy check enforces it.
"""
from __future__ import annotations

import hashlib
import io
import math
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.models.db import KnowledgeChunk, KnowledgeDoc
from app.services.state_service import record_event

ALLOWED_EXT = {".txt", ".md", ".markdown", ".docx", ".pdf", ".html", ".htm", ".csv", ".json", ".xml"}
CHUNK_CHARS = 900
MIN_SCORE = 1.0          # below this a passage is not considered relevant
_STOP = set("""a an and are as at be by for from has have if in into is it its of on or that the this to was were
will with not no can may must should use using used when than then there their they we you your our all any
each per via also only more most such""".split())
_TOKEN = re.compile(r"[a-z0-9][a-z0-9_.:-]*[a-z0-9]|[a-z0-9]")


class KnowledgeError(Exception):
    pass


@dataclass
class Hit:
    chunk_id: int
    ref: str
    doc_id: int
    doc_title: str
    heading: str | None
    text: str
    score: float

    def as_context(self) -> dict:
        src = self.doc_title + (f" > {self.heading}" if self.heading else "")
        return {"id": self.ref, "source": src, "text": self.text}


# --- text extraction -----------------------------------------------------------
def extract_sections(filename: str, data: bytes) -> list[tuple[str | None, str]]:
    """Return [(heading, text)] blocks in document order."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise KnowledgeError(f"Unsupported file type {ext or '(none)'}; use {', '.join(sorted(ALLOWED_EXT))}")
    if ext == ".docx":
        import zipfile
        import docx
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:  # refuse zip bombs before python-docx unpacks it
                if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                    raise KnowledgeError("The .docx expands to more than 100 MB")
            d = docx.Document(io.BytesIO(data))
        except KnowledgeError:
            raise
        except Exception as exc:  # noqa: BLE001 - corrupt or not really a .docx
            raise KnowledgeError(f"Could not read the .docx ({type(exc).__name__})") from exc
        out, heading = [], None
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            if (p.style.name or "").lower().startswith(("heading", "title")):
                heading = t
                continue
            out.append((heading, t))
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    out.append((heading, " | ".join(dict.fromkeys(cells))))
        return out
    if ext == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data))
            return [(f"page {i}", (pg.extract_text() or "").strip()) for i, pg in enumerate(reader.pages, 1)
                    if (pg.extract_text() or "").strip()]
        except Exception as exc:  # corrupt / encrypted PDF
            raise KnowledgeError(f"Could not read the PDF: {exc}") from exc
    text = data.decode("utf-8", errors="replace")
    if ext in (".html", ".htm"):
        text = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", text)
        text = re.sub(r"(?i)<h[1-6][^>]*>", "\n# ", text)
        text = re.sub(r"(?i)<(br|/p|/div|/li|/tr|/h[1-6])[^>]*>", "\n", text)
        text = re.sub(r"<[^>]+>", " ", text)
        import html
        text = html.unescape(text)
    out, heading, para = [], None, []

    def flush():
        if para:
            out.append((heading, " ".join(para).strip()))
            para.clear()
    for line in text.splitlines():
        s = line.strip()
        m = re.match(r"^#{1,6}\s+(.*)", s)
        if m:
            flush()
            heading = m.group(1).strip()[:300]
        elif not s:
            flush()
        else:
            para.append(s)
    flush()
    return [(h, t) for h, t in out if t]


def chunk(sections: list[tuple[str | None, str]], size: int = CHUNK_CHARS) -> list[tuple[str | None, str]]:
    """Merge paragraphs under the same heading up to ~size characters; split long ones."""
    chunks: list[tuple[str | None, str]] = []
    cur_h, cur = None, ""
    for h, t in sections:
        pieces = [t] if len(t) <= size else [t[i:i + size] for i in range(0, len(t), size - 100)]
        for p in pieces:
            if cur and (h != cur_h or len(cur) + len(p) + 1 > size):
                chunks.append((cur_h, cur))
                cur = ""
            cur_h = h
            cur = f"{cur}\n{p}".strip()
    if cur:
        chunks.append((cur_h, cur))
    return chunks


# --- documents ----------------------------------------------------------------
def add_document(session: Session, filename: str, data: bytes, actor: str, title: str | None = None,
                 app_id: str | None = None, tags: list[str] | None = None, max_mb: int = 10) -> KnowledgeDoc:
    actor = (actor or "").strip()
    if not actor:
        raise KnowledgeError("Your name is required to upload")
    if not data:
        raise KnowledgeError("The file is empty")
    if len(data) > max_mb * 1024 * 1024:
        raise KnowledgeError(f"File is larger than {max_mb} MB")
    digest = hashlib.sha256(data).hexdigest()
    dup = session.scalars(select(KnowledgeDoc).where(KnowledgeDoc.sha256 == digest)).first()
    if dup:
        raise KnowledgeError(f"Already uploaded as '{dup.title}'")
    parts = chunk(extract_sections(filename, data))
    if not parts:
        raise KnowledgeError("No text found in the file (scanned PDFs need OCR first)")
    doc = KnowledgeDoc(title=(title or Path(filename).stem)[:500], filename=Path(filename).name[:500], sha256=digest,
                       app_id=app_id or None, tags=[t.strip() for t in (tags or []) if t.strip()][:10],
                       uploaded_by=actor, chars=sum(len(t) for _, t in parts))
    doc.chunks = [KnowledgeChunk(position=i, heading=h, text=t) for i, (h, t) in enumerate(parts)]
    session.add(doc)
    session.flush()
    record_event(session, "KNOWLEDGE_ADDED", actor, app_id or None,
                 {"doc_id": doc.id, "title": doc.title, "chunks": len(parts), "sha256": digest})
    return doc


def delete_document(session: Session, doc: KnowledgeDoc, actor: str) -> None:
    actor = (actor or "").strip()
    if not actor:
        raise KnowledgeError("Your name is required to delete")
    record_event(session, "KNOWLEDGE_DELETED", actor, doc.app_id,
                 {"doc_id": doc.id, "title": doc.title, "sha256": doc.sha256})
    session.delete(doc)


# --- search (BM25) --------------------------------------------------------------
def tokenize(text: str) -> list[str]:
    toks = []
    for t in _TOKEN.findall((text or "").lower()):
        t = t.strip(".:-")
        if not t or t in _STOP or len(t) == 1:
            continue
        if len(t) > 4 and t.endswith("s") and not t.endswith("ss"):
            t = t[:-1]
        toks.append(t)
        if any(c in t for c in ".:-_"):          # also index the parts of urn:..., user.email, entity-id
            toks += [p for p in re.split(r"[.:_-]+", t) if len(p) > 1 and p not in _STOP]
    return toks


def _candidates(session: Session, app_id: str | None) -> list[KnowledgeChunk]:
    q = select(KnowledgeChunk).join(KnowledgeDoc)
    q = q.where(or_(KnowledgeDoc.app_id.is_(None), KnowledgeDoc.app_id == app_id)) if app_id else \
        q.where(KnowledgeDoc.app_id.is_(None))
    return list(session.scalars(q))


def search(session: Session, query: str, k: int = 5, app_id: str | None = None,
           min_score: float = MIN_SCORE, _cache: dict | None = None) -> list[Hit]:
    """BM25 (k1=1.5, b=0.75) over the chunks visible to this app (global + app-specific).
    App-specific documents get a small boost."""
    q = [t for t in dict.fromkeys(tokenize(query))]
    if not q:
        return []
    if _cache is not None and "index" in _cache:
        chunks, docs_tf, df, avgdl = _cache["index"]
    else:
        chunks = _candidates(session, app_id)
        if not chunks:
            return []
        docs_tf = [Counter(tokenize((c.heading or "") + " " + c.text)) for c in chunks]
        df = Counter(t for tf in docs_tf for t in tf)
        avgdl = sum(sum(tf.values()) for tf in docs_tf) / len(chunks) or 1
        if _cache is not None:
            _cache["index"] = (chunks, docs_tf, df, avgdl)
    n = len(chunks)
    scored = []
    for c, tf in zip(chunks, docs_tf):
        dl = sum(tf.values()) or 1
        s = 0.0
        for t in q:
            if t in tf:
                idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
                s += idf * tf[t] * 2.5 / (tf[t] + 1.5 * (1 - 0.75 + 0.75 * dl / avgdl))
        if s and c.doc.app_id and c.doc.app_id == app_id:
            s *= 1.25
        if s >= min_score:
            scored.append((s, c))
    scored.sort(key=lambda x: (-x[0], x[1].id))
    return [Hit(c.id, c.ref, c.doc_id, c.doc.title, c.heading, c.text, round(s, 2)) for s, c in scored[:k]]


def stats(session: Session) -> dict:
    docs = session.scalars(select(KnowledgeDoc)).all()
    return {"docs": len(docs), "chunks": sum(len(d.chunks) for d in docs)}
