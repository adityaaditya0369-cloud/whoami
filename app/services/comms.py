"""App-owner / vendor communication pack: one ready-to-send message per app.

Never includes secrets. New OIDC client secrets are delivered through the
channel chosen in the decisions register, never by email.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import Settings
from app.models.db import Application
from app.services import comparison
from app.services import estimate as em


@dataclass
class Message:
    app: Application
    to: str
    subject: str
    body: str
    pilot: object
    cutover: object


def _dates(app: Application, tl: em.Timeline):
    wave = app.wave or app.suggested_wave
    cut = tl.cutover.get(wave)
    pilot = cut - timedelta(days=tl.settings.pilot_days + 1) if cut else None
    return wave, pilot, cut


def build_message(app: Application, settings: Settings, tl: em.Timeline, questions: list[str],
                  secret_channel: str | None = None) -> Message:
    wave, pilot, cut = _dates(app, tl)
    rows = comparison.build(app, settings)
    new_values = [r for r in rows if r.status == "SP_UPDATE"]
    unknown = [r for r in rows if r.status == "UNKNOWN"]
    when = f"{cut:%d %b %Y}" if cut else "to be agreed"
    subject = f"[Action needed] Sign-in for {app.label} moves from Okta to PingFederate - planned {when}"
    L = []
    L.append(f"Hello{(' ' + app.business_owner) if app.business_owner else ''},")
    L.append("")
    L.append(f"As part of moving single sign-on from Okta to PingFederate, {app.label} is planned for "
             f"{wave or 'a later wave'}.")
    L.append("")
    L.append("PROPOSED DATES")
    L.append(f"  Pilot with a few test users: from {pilot:%d %b %Y}" if pilot else "  Pilot: to be agreed")
    L.append(f"  Production cutover: {when}")
    L.append("  Okta stays in place until the cutover is verified, so we can roll back quickly if needed.")
    L.append("")
    if app.is_saml:
        L.append("WHAT CHANGES ON YOUR SIDE (SAML)")
        for r in new_values:
            if r.field in ("IdP-initiated SSO link",):
                continue
            L.append(f"  {r.field}: {r.pf_value}")
        L.append("  Please configure the new identity provider alongside Okta for the pilot, if your application allows it.")
        L.append("")
        L.append("WHAT STAYS THE SAME")
        same = [r for r in rows if r.status == "SAME" and r.section in ("SP connection", "Assertion", "Attributes")]
        for r in same[:8]:
            L.append(f"  {r.field}: {r.okta_value}")
    else:
        o = app.oidc
        L.append("WHAT CHANGES ON YOUR SIDE (OpenID Connect)")
        for r in new_values:
            if r.section == "Endpoints":
                L.append(f"  {r.field}: {r.pf_value}")
        L.append(f"  Client ID stays the same: {o.client_id}")
        if o.token_endpoint_auth_method in ("client_secret_basic", "client_secret_post", "client_secret_jwt"):
            L.append(f"  A NEW client secret will be issued. It will be shared via {secret_channel or 'a secure channel'}, "
                     "never by email.")
        if "refresh_token" in (o.grant_types or []):
            L.append("  Users will sign in once more after the cutover (existing refresh tokens cannot be moved).")
        if "implicit" in (o.grant_types or []):
            L.append("  The implicit grant is deprecated; we propose switching to authorization code + PKCE.")
    L.append("")
    asks = list(questions)
    if unknown and app.is_saml and app.saml and app.saml.config_completeness == "PARTIAL":
        asks.insert(0, "Please send your SAML SP metadata (entity ID, ACS URLs, required attributes).")
    if app.has_test_environment is None:
        asks.append("Is there a test or sandbox instance we can switch first?")
    if asks:
        L.append("WHAT WE NEED FROM YOU")
        for i, q in enumerate(dict.fromkeys(asks), 1):
            L.append(f"  {i}. {q}")
        L.append("")
    L.append("Please reply to confirm the dates, or suggest a better change window.")
    L.append("")
    L.append("Thank you,")
    L.append("Identity migration team")
    return Message(app, app.business_owner or "", subject, "\n".join(L), pilot, cut)


def to_xlsx(messages: list[Message]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mail merge"
    headers = ["Application", "Protocol", "To (owner / vendor)", "Subject", "Body", "Pilot from", "Cutover", "Sent?"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2A44")
    for r, m in enumerate(messages, 2):
        vals = [m.app.label, m.app.protocol, m.to, m.subject, m.body,
                m.pilot.date() if m.pilot else None, m.cutover.date() if m.cutover else None, ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=i in (4, 5))
        ws.cell(row=r, column=6).number_format = ws.cell(row=r, column=7).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=8).fill = PatternFill("solid", fgColor="FFFF00")
    for col, w in zip("ABCDEFGH", (28, 9, 26, 50, 90, 12, 12, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
