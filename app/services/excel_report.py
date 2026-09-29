"""Excel migration report (openpyxl). One workbook for all apps or a single app."""
from __future__ import annotations

import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.config import Settings
from app.models.db import Application, utcnow
from app.services import comparison, risk
from app.services.ai.service import latest_assessment
from app.services import plan as planmod
from app.services.readiness import pf_readiness
from app.services.reconcile import BUILD_LABEL, BUILD_STATUSES, build_status, reconcile

FONT = "Arial"
NAVY = "1F2A44"
THIN = Side(style="thin", color="D9DDE5")
STATUS_FILL = {
    "ACTION": "FDE2E1", "UNKNOWN": "FFF1D6", "SP_UPDATE": "E4ECFD", "CONFIGURE": "EEE8FB",
    "SAME": "E3F4EA", "N/A": "F2F3F5",
}
LEVEL_FILL = {"LOW": "E3F4EA", "MEDIUM": "FFF1D6", "HIGH": "FDE2E1", "CRITICAL": "F4B6B2"}
SAML_SHEET = "SAML - Okta vs PingFederate"
OIDC_SHEET = "OIDC - Okta vs PingFederate"
SEV_FILL = {"CRITICAL": "F4B6B2", "WARNING": "FFF1D6", "INFO": "E4ECFD"}


def _header(ws: Worksheet, headers: list[str], widths: list[int], row: int = 1) -> None:
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _rows(ws: Worksheet, rows: list[list], start: int = 2, wrap_cols: set[int] | None = None) -> None:
    wrap_cols = wrap_cols or set()
    for r_i, row in enumerate(rows, start=start):
        for c_i, v in enumerate(row, start=1):
            if isinstance(v, str) and v[:1] in "=+-@" and not v.startswith("=COUNT"):
                v = "'" + v  # formula-injection guard for data values
            c = ws.cell(row=r_i, column=c_i, value=v)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=c_i in wrap_cols)
            c.border = Border(bottom=THIN)
    if rows:
        ws.auto_filter.ref = f"A{start - 1}:{get_column_letter(len(rows[0]))}{start - 1 + len(rows)}"


def _color_col(ws: Worksheet, col: int, fills: dict[str, str], first: int, last: int) -> None:
    if last < first:
        return
    letter = get_column_letter(col)
    rng = f"{letter}{first}:{letter}{last}"
    for value, color in fills.items():
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'${letter}{first}="{value}"'], fill=PatternFill("solid", fgColor=color),
            font=Font(name=FONT, bold=True, size=10)))


def _readiness(app: Application) -> dict:
    return pf_readiness(app)


def build_workbook(session, apps: list[Application], settings: Settings, title: str) -> bytes:
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_apps = wb.create_sheet("Applications")
    ws_cmp = wb.create_sheet(SAML_SHEET)
    ws_ocmp = wb.create_sheet(OIDC_SHEET)
    ws_recon = wb.create_sheet("PingFederate reconciliation")
    ws_plan2 = wb.create_sheet("Migration plan")
    ws_pipe = wb.create_sheet("Pipeline")
    ws_val = wb.create_sheet("SAML validation")
    ws_agents = wb.create_sheet("Agent reviews")
    ws_attr = wb.create_sheet("Attribute mapping")
    ws_prep = wb.create_sheet("Directory prep")
    ws_plan = wb.create_sheet("Action plan")
    ws_q = wb.create_sheet("Owner questions")
    ws_find = wb.create_sheet("Findings")
    ws_risk = wb.create_sheet("Risk breakdown")

    # --- Applications ---------------------------------------------------------
    headers = ["Okta app ID", "Application", "Protocol", "Type", "Okta status", "Migration state", "Users",
               "Direct users", "Groups", "Claims", "Complexity score", "Complexity", "Impact score",
               "Impact", "Overall", "Blocked", "Blockers", "Suggested wave", "Assigned wave",
               "Business owner", "Criticality", "Test env", "Issuance criteria needed", "OGNL values",
               "Okta-only groups", "LDAP attributes", "Missing directory inputs", "Critical", "Warning",
               "Info", "Assessment", "Assessment review", "Reviewed by", "PingFederate build",
               "Plan steps done", "Next plan step"]
    _header(ws_apps, headers, [22, 32, 10, 16, 11, 16, 8, 8, 8, 8, 10, 12, 9, 11, 11, 9, 30, 20, 14,
                               20, 11, 9, 12, 24, 26, 30, 24, 8, 8, 8, 14, 14, 16, 20, 10, 26])
    ocmp_rows: list[list] = []
    app_rows, cmp_rows, attr_rows, find_rows, risk_rows, plan_rows, q_rows = [], [], [], [], [], [], []
    prep: dict[tuple[str, str], set[str]] = defaultdict(set)
    recon_by_app = {r.app.id: r for r in reconcile(session)[0]}
    recon_rows, plan_rows2 = [], []
    for a in apps:
        rc = recon_by_app.get(a.id)
        tasks = planmod.ensure_tasks(session, a) if a.state != "OUT_OF_SCOPE" else {}
        bstat = build_status(a, rc, tasks)
        done, total = planmod.progress(tasks) if tasks else (0, 0)
        nxt = planmod.next_step(tasks) if tasks else None
        if rc and rc.pf:
            for c in rc.checks:
                recon_rows.append([a.label, a.protocol, rc.pf.name, rc.matched_by, c.check, c.expected, c.actual,
                                   "OK" if c.ok else "DIFFERS"])
        else:
            recon_rows.append([a.label, a.protocol, "", "", "Exists on PingFederate", "Yes", "Not found", "NOT BUILT"])
        for i, step in enumerate(planmod.STEPS, start=1):
            t = tasks.get(step.key)
            if t is not None:
                plan_rows2.append([a.label, a.protocol, planmod.effective_wave(a), i, step.title,
                                   t.status.replace("_", " ").title(), t.owner or step.owner,
                                   t.due_date.date() if t.due_date else None, t.notes or "",
                                   t.updated_by or "", BUILD_LABEL[bstat]])
        pf = _readiness(a)
        score = risk.latest_score(session, a.id)
        assess = latest_assessment(session, a.id)
        sev = defaultdict(int)
        for f in a.findings:
            sev[f.severity] += 1
        app_rows.append([
            a.id, a.label, a.protocol,
            ("Custom" if a.is_custom_saml else "Catalog") if a.is_saml else ((a.oidc.application_type or "").title() if a.oidc else ""),
            a.okta_status, a.state,
            a.user_count, a.direct_user_count, a.group_count, len(a.claims),
            a.complexity_score, a.complexity_level, a.impact_score, a.impact_level, a.overall_level,
            "Yes" if a.blocked else "No", ", ".join(score.blockers) if score else "",
            a.suggested_wave or "", a.wave or "", a.business_owner or "", a.business_criticality or "",
            {True: "Yes", False: "No"}.get(a.has_test_environment, "Unknown"),
            "Yes" if (pf["issuance_criteria"] or any(f.code == "OIDC_ACCESS_CONTROL_REQUIRED" for f in a.findings)) else "No", ", ".join(pf["ognl_claims"]),
            ", ".join(pf["okta_only_groups"]), ", ".join(pf["ldap_attributes"]),
            ", ".join(pf["missing_inputs"]), sev["CRITICAL"], sev["WARNING"], sev["INFO"],
            assess.status if assess else "Not run", assess.review_status if assess else "",
            (assess.reviewed_by or "") if assess else "",
            BUILD_LABEL[bstat], f"{done}/{total}" if total else "", nxt.title if nxt else ("All done" if total else ""),
        ])
        for r in comparison.build(a, settings):
            (ocmp_rows if a.is_oidc else cmp_rows).append([a.label, r.section, r.field, r.okta_value, r.pf_setting, r.pf_value, r.status, r.note])
        for c in a.claims:
            attr_rows.append([a.label, c.name, c.claim_type,
                              f"{c.group_filter_type} {c.group_filter_value}" if c.claim_type == "GROUP" else ", ".join(c.values),
                              c.expression_kind or "GROUP", c.pf_source, c.pf_source_detail,
                              ", ".join(c.pf_ldap_attributes or []), ", ".join(c.pf_missing_inputs or []),
                              ", ".join(c.matched_okta_native_groups or [])])
        for f in sorted(a.findings, key=lambda f: ({"CRITICAL": 0, "WARNING": 1, "INFO": 2}[f.severity], f.code)):
            find_rows.append([a.label, f.severity, f.code, f.message])
        if score:
            for b in score.breakdown:
                risk_rows.append([a.label, b["dimension"], b["rule"], b["points"], b["reason"], score.rules_version])
        for g in pf["okta_only_groups"]:
            prep[("Group to create", g)].add(a.label)
        for x in pf["ldap_attributes"]:
            prep[("LDAP attribute read", x)].add(a.label)
        for x in pf["missing_inputs"]:
            prep[("Value with no directory source", x)].add(a.label)
        if assess and assess.status == "VALID":
            src = f"{assess.provider}" + (f" ({assess.model})" if assess.model else "")
            for act in assess.output.get("recommended_actions", []):
                plan_rows.append([a.label, act["when"].replace("_", " ").title(), act["owner"].replace("_", " ").title().replace("Iam ", "IAM "),
                                  act["action"], ", ".join(act.get("related_codes", [])), src, assess.review_status])
            for q in assess.output.get("questions_for_app_owner", []):
                q_rows.append([a.label, a.business_owner or "Unknown", q, ""])

    _rows(ws_apps, app_rows, wrap_cols={16, 23, 24, 25, 26})
    n = len(app_rows) + 1
    for col, fills in ((headers.index("Complexity") + 1, LEVEL_FILL), (headers.index("Impact") + 1, LEVEL_FILL),
                       (headers.index("Overall") + 1, LEVEL_FILL)):
        _color_col(ws_apps, col, fills, 2, n)
    _color_col(ws_apps, headers.index("Blocked") + 1, {"Yes": "F4B6B2"}, 2, n)

    for ws_c, rows_c in ((ws_cmp, cmp_rows), (ws_ocmp, ocmp_rows)):
        _header(ws_c, ["Application", "Section", "Field", "Okta value", "PingFederate setting",
                       "PingFederate value", "Status", "Note"], [28, 18, 26, 44, 40, 44, 12, 40])
        _rows(ws_c, rows_c, wrap_cols={4, 5, 6, 8})
        _color_col(ws_c, 7, STATUS_FILL, 2, len(rows_c) + 1)

    _header(ws_recon, ["Okta app", "Protocol", "PingFederate object", "Matched by", "Check", "Expected (Okta)",
                       "Actual (PingFederate)", "Result"], [30, 9, 30, 12, 30, 44, 44, 12])
    _rows(ws_recon, recon_rows, wrap_cols={6, 7})
    _color_col(ws_recon, 8, {"OK": "E3F4EA", "DIFFERS": "FDE2E1", "NOT BUILT": "E4ECFD"}, 2, len(recon_rows) + 1)

    _header(ws_plan2, ["Application", "Protocol", "Wave", "#", "Step", "Status", "Owner", "Due date", "Notes",
                       "Updated by", "PingFederate build"], [30, 9, 20, 5, 30, 13, 22, 12, 40, 16, 20])
    _rows(ws_plan2, plan_rows2, wrap_cols={9})
    for r_ in range(2, len(plan_rows2) + 2):
        ws_plan2.cell(row=r_, column=8).number_format = "yyyy-mm-dd"
    _color_col(ws_plan2, 6, {"Done": "E3F4EA", "In Progress": "E4ECFD", "Blocked": "FDE2E1"}, 2, len(plan_rows2) + 1)

    _header(ws_attr, ["Application", "Attribute", "Type", "Okta value / filter", "Okta classification",
                      "PingFederate source", "PingFederate detail", "LDAP attributes", "Missing inputs",
                      "Okta-only groups"], [28, 18, 12, 40, 14, 18, 50, 22, 22, 24])
    _rows(ws_attr, attr_rows, wrap_cols={4, 7})
    _color_col(ws_attr, 6, {"DATA_STORE": "E3F4EA", "TEXT": "E3F4EA", "GROUP_LDAP_SEARCH": "EEE8FB",
                            "OGNL": "FDE2E1", "GROUP_OGNL": "FDE2E1", "APPUSER": "FDE2E1",
                            "UNMAPPED": "FDE2E1"}, 2, len(attr_rows) + 1)

    prep_rows = [[k[0], k[1], len(v), ", ".join(sorted(v))] for k, v in sorted(prep.items())]
    _header(ws_prep, ["Item", "Name", "Apps affected", "Applications"], [30, 30, 12, 60])
    _rows(ws_prep, prep_rows, wrap_cols={4})

    _header(ws_plan, ["Application", "When", "Owner", "Action", "Resolves findings", "Source", "Review"],
            [28, 18, 16, 70, 34, 26, 12])
    _rows(ws_plan, plan_rows, wrap_cols={4, 5})

    _header(ws_q, ["Application", "Business owner", "Question", "Answer (fill in)"], [28, 22, 70, 50])
    _rows(ws_q, q_rows, wrap_cols={3, 4})
    for r in range(2, len(q_rows) + 2):
        ws_q.cell(row=r, column=4).fill = PatternFill("solid", fgColor="FFFF00")

    _header(ws_find, ["Application", "Severity", "Code", "Message"], [28, 11, 32, 100])
    _rows(ws_find, find_rows, wrap_cols={4})
    _color_col(ws_find, 2, SEV_FILL, 2, len(find_rows) + 1)

    _header(ws_risk, ["Application", "Dimension", "Rule", "Points", "Reason", "Rules version"],
            [28, 13, 32, 8, 50, 16])
    _rows(ws_risk, risk_rows, wrap_cols={5})

    # --- Pipeline / validation / agents ----------------------------------------------
    from sqlalchemy import select as _select

    from app.models.db import AgentRun, MigrationPlan, ValidationResult
    from app.services import pipeline as pipemod
    from app.models.state import MigrationState
    from app.services.agents.schema import AGENT_TITLE
    stage_title = {k: t for k, t, _ in pipemod.STAGES}
    ids = [a.id for a in apps if a.is_saml]
    runs = {r.app_id: r for r in session.scalars(_select(AgentRun).where(AgentRun.app_id.in_(ids)).order_by(AgentRun.id))}
    plans = {p.app_id: p for p in session.scalars(_select(MigrationPlan).where(MigrationPlan.app_id.in_(ids)).order_by(MigrationPlan.id))}
    vals = list(session.scalars(_select(ValidationResult).where(ValidationResult.app_id.in_(ids)).order_by(ValidationResult.id)))
    last_val = {}
    for v in vals:
        last_val[(v.app_id, v.stage)] = v
    label = {a.id: a.label for a in apps}
    pipe_rows, val_rows, agent_rows = [], [], []
    for a in apps:
        if not a.is_saml or a.state == "OUT_OF_SCOPE":
            continue
        run, pl_ = runs.get(a.id), plans.get(a.id)
        cnt = {"AGREE": 0, "CONCERN": 0, "DISAGREE": 0}
        if run:
            for rv in run.reviews:
                for it in (rv.output or {}).get("items", []) if rv.status == "VALID" else []:
                    cnt[it.get("assessment", "AGREE")] += 1
                    if it.get("assessment") != "AGREE":
                        agent_rows.append([a.label, AGENT_TITLE.get(rv.agent, rv.agent), it.get("key"), it.get("assessment"),
                                           it.get("comment"), it.get("suggestion") or "", ", ".join(it.get("citations") or []),
                                           run.review_status, (run.resolutions or {}).get(f"{rv.agent}:{it.get('key')}", "")])
        pre, post = last_val.get((a.id, "PRE_CUTOVER")), last_val.get((a.id, "POST_CUTOVER"))
        base = last_val.get((a.id, "BASELINE_OKTA"))
        pipe_rows.append([a.label, stage_title[pipemod.stage_of(a.state)], a.state.replace("_", " ").title(),
                          (run.review_status.title() if run and run.status == "VALID" else (run.status.title() if run else "")),
                          cnt["DISAGREE"] if run else "", cnt["CONCERN"] if run else "",
                          f"v{pl_.version}" if pl_ else "", pl_.status.title() if pl_ else "",
                          pl_.decided_by if pl_ and pl_.status == "APPROVED" else "",
                          base.verdict.replace("_", " ").title() if base else "",
                          pre.verdict.replace("_", " ").title() if pre else "",
                          post.verdict.replace("_", " ").title() if post else "",
                          pipemod.NEXT_ACTION.get(MigrationState(a.state), "")])
    for v in reversed(vals):
        fails = [c["label"] for c in v.result.get("checks", []) if c["status"] == "FAIL"]
        warns = [c["label"] for c in v.result.get("checks", []) if c["status"] == "WARN"]
        val_rows.append([label.get(v.app_id, v.app_id), v.id, v.created_at.strftime("%Y-%m-%d %H:%M"),
                         v.stage.replace("_", " ").title(), (v.source or "").title(), v.verdict.replace("_", " ").title(),
                         "; ".join(fails), "; ".join(warns), v.actor or "", "Yes" if v.result.get("baseline_used") else "No"])
    _header(ws_pipe, ["Application", "Stage", "State", "Agent review", "Disagreements", "Concerns", "Plan", "Plan status",
                      "Approved by", "Okta baseline", "Pre-cutover", "Post-cutover", "Next step"],
            [30, 16, 20, 13, 13, 10, 7, 12, 14, 14, 14, 14, 44])
    _rows(ws_pipe, pipe_rows, wrap_cols={13})
    vfill = {"Pass": "E3F4EA", "Pass With Warnings": "FFF1D6", "Fail": "FDE2E1"}
    for c in (10, 11, 12):
        _color_col(ws_pipe, c, vfill, 2, len(pipe_rows) + 1)
    _header(ws_val, ["Application", "#", "When (UTC)", "Stage", "Source", "Verdict", "Failed checks", "Warnings", "By",
                     "Compared with Okta"], [30, 6, 17, 16, 11, 18, 50, 40, 14, 12])
    _rows(ws_val, val_rows, wrap_cols={7, 8})
    _color_col(ws_val, 6, vfill, 2, len(val_rows) + 1)
    _header(ws_agents, ["Application", "Agent", "Item", "Assessment", "Comment", "Suggestion", "Sources", "Review",
                        "Resolution"], [28, 22, 26, 12, 60, 44, 14, 11, 40])
    _rows(ws_agents, agent_rows, wrap_cols={5, 6, 9})
    _color_col(ws_agents, 4, {"DISAGREE": "FDE2E1", "CONCERN": "FFF1D6"}, 2, len(agent_rows) + 1)

    # --- Summary (formulas over the other sheets) -------------------------------
    ws = ws_sum
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=NAVY)
    meta = [("Customer", settings.customer_name), ("Generated (UTC)", utcnow().strftime("%Y-%m-%d %H:%M")),
            ("Risk rules version", risk.RULES_VERSION),
            ("Target", "PingFederate (self-hosted)"), ("Directory", settings.pf_directory_type.value),
            ("OGNL allowed", "Yes" if settings.pf_ognl_allowed else "No"),
            ("PingFederate base URL (assumed)", settings.pf_base_url),
            ("PingFederate entity ID (assumed)", settings.pf_entity_id)]
    r = 3
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        ws.cell(row=r, column=2, value=v).font = Font(name=FONT, size=10)
        r += 1
    ws.cell(row=r, column=3, value="Base URL and entity ID come from PF_BASE_URL / PF_ENTITY_ID in .env").font = \
        Font(name=FONT, italic=True, size=9, color="666666")

    def block(start: int, heading: str, items: list[tuple], cols: tuple[str, ...] = ()) -> int:
        ws.cell(row=start, column=1, value=heading).font = Font(name=FONT, bold=True, size=12, color=NAVY)
        for j, h in enumerate(cols, start=2):
            c = ws.cell(row=start, column=j, value=h)
            c.font = Font(name=FONT, bold=True, size=10, color="666666")
            c.alignment = Alignment(horizontal="right")
        rr = start + 1
        for label, *formulas in items:
            ws.cell(row=rr, column=1, value=label).font = Font(name=FONT, size=10)
            for j, formula in enumerate(formulas, start=2):
                c = ws.cell(row=rr, column=j, value=formula)
                c.font = Font(name=FONT, size=10, bold=True)
                c.alignment = Alignment(horizontal="right")
            rr += 1
        return rr + 1

    def col(name: str) -> str:
        return get_column_letter(headers.index(name) + 1)

    last = max(n, 2)
    A = "Applications"
    P = f"{A}!${col('Protocol')}$2:${col('Protocol')}${last}"

    def by_proto(column: str, value: str) -> tuple[str, str]:
        rng = f"{A}!${col(column)}$2:${col(column)}${last}"
        return (f'=COUNTIFS({P},"SAML",{rng},"{value}")', f'=COUNTIFS({P},"OIDC",{rng},"{value}")')

    users = f"{A}!${col('Users')}$2:${col('Users')}${last}"
    r = block(r + 2, "Applications", [
        ("Applications", f'=COUNTIF({P},"SAML")', f'=COUNTIF({P},"OIDC")'),
        ("Blocked", *by_proto("Blocked", "Yes")),
        ("Need access control (issuance criteria / policy)", *by_proto("Issuance criteria needed", "Yes")),
        ("Total assigned users", f'=SUMIF({P},"SAML",{users})', f'=SUMIF({P},"OIDC",{users})'),
    ], cols=("SAML", "OIDC"))
    r = block(r, "Overall level", [(lvl, *by_proto("Overall", lvl)) for lvl in risk.LEVELS], cols=("SAML", "OIDC"))
    waves = ["Wave 0 (pilot)", "Wave 1", "Wave 2", "Wave 3 (high impact)", "Blocked", "Decommission review"]
    r = block(r, "Suggested wave", [(w, *by_proto("Suggested wave", w)) for w in waves], cols=("SAML", "OIDC"))
    r = block(r, "PingFederate build status", [(BUILD_LABEL[b], *by_proto("PingFederate build", BUILD_LABEL[b]))
                                               for b in BUILD_STATUSES], cols=("SAML", "OIDC"))
    pl = max(len(plan_rows2) + 1, 2)
    r = block(r, "Migration plan steps", [
        (st.replace("_", " ").title(), f"=COUNTIFS('Migration plan'!$B$2:$B${pl},\"SAML\",'Migration plan'!$F$2:$F${pl},\"{st.replace('_', ' ').title()}\")",
         f"=COUNTIFS('Migration plan'!$B$2:$B${pl},\"OIDC\",'Migration plan'!$F$2:$F${pl},\"{st.replace('_', ' ').title()}\")")
        for st in planmod.STEP_STATUSES], cols=("SAML", "OIDC"))
    ppl = max(len(pipe_rows) + 1, 2)
    r = block(r, "SAML pipeline stage", [(t, f"=COUNTIF(Pipeline!$B$2:$B${ppl},\"{t}\")")
                                         for _, t, _ in pipemod.STAGES], cols=("SAML",))
    cl, ol = max(len(cmp_rows) + 1, 2), max(len(ocmp_rows) + 1, 2)
    r = block(r, "Okta vs PingFederate fields", [
        (f"{st} - {comparison.STATUS_HELP[st]}",
         f"=COUNTIF('{SAML_SHEET}'!$G$2:$G${cl},\"{st}\")", f"=COUNTIF('{OIDC_SHEET}'!$G$2:$G${ol},\"{st}\")")
        for st in comparison.STATUS_ORDER], cols=("SAML", "OIDC"))
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = sheet.title != "Summary"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
