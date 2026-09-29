"""Migration Strategy Pack as Excel (all views) and Word (executive + architecture narrative).
Every figure comes from the pack; nothing here computes a new number."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.strategy_model.service import STRATEGY_HELP, StrategyPack

FONT = "Calibri"
NAVY = "1F2A44"
THIN = Side(style="thin", color="D9DDE5")
FILL = {"RECREATE": "E6F4EA", "TRANSFORM": "E8F0FE", "REDESIGN": "F3E8FD", "RETIRE": "EEEEEE", "RETAIN": "FEF7E0",
        "EQUIVALENT": "E6F4EA", "PARTIAL": "FEF7E0", "NONE": "FCE8E6", "HIGH": "E6F4EA", "MEDIUM": "FEF7E0", "LOW": "FCE8E6"}


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], rows: list[list],
           color_cols: tuple[int, ...] = ()) -> None:
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, start=2):
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=THIN)
            if i in color_cols and isinstance(v, str) and v in FILL:
                c.fill = PatternFill("solid", fgColor=FILL[v])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"


def _d(v):
    return v.strftime("%d %b %Y") if v else ""


def to_xlsx(p: StrategyPack, customer: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive summary"
    s = p.summary
    ws["A1"] = f"Migration Strategy Pack: Okta → {p.target}"
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=NAVY)
    ws["A2"] = f"{customer} · generated {p.generated:%d %b %Y %H:%M} UTC · catalog {p.catalog_version} · risk rules {p.rules_version}"
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="666666")
    items = [("Applications in Okta", s["total"]), ("Federated (SAML + OIDC)", s["federated"]),
             ("Potentially migratable (recreate + transform)", s["migratable"]), ("  Recreate", s["recreate"]),
             ("  Transform", s["transform"]), ("Requires redesign", s["redesign"]), ("Retire", s["retire"]),
             ("Retain on Okta for now", s["retain"]), ("High-risk applications", s["high_risk"]),
             ("Migration waves", s["waves"]), ("Tenant objects analysed", s["tenant_objects"]),
             ("Tenant objects needing redesign", s["tenant_redesign"]), ("Foundation work (h)", s["foundation_hours"]),
             ("Application work (h)", s["app_hours"]), ("Planned finish", _d(s["end"])),
             ("Confirmed dependencies", s["deps_confirmed"]), ("Dependencies awaiting confirmation", s["deps_suggested"]),
             ("Catalog rows needing expert review", s["expert_rows"])]
    for i, (k, v) in enumerate(items, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(name=FONT, size=10, bold=not k.startswith("  "))
        ws.cell(row=i, column=2, value=v).font = Font(name=FONT, size=10)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 18

    _sheet(wb, "Portfolio (architect)",
           ["Application", "Protocol", "Strategy", "Compatibility", "Complexity", "Risk", "Wave", "Wave notes",
            "Depends on", "Used by", "Owner", "Tests", "Effort (h)", "Main reason"],
           [34, 9, 12, 13, 11, 10, 20, 40, 30, 30, 20, 7, 10, 60],
           [[r.app.label, r.app.protocol, r.strategy, r.compat, r.complexity, r.risk, r.wave,
             "; ".join(r.wave_reasons + r.flags), ", ".join(f"{a} ({k})" for a, k in r.depends_on),
             ", ".join(f"{a} ({k})" for a, k in r.depended_by), r.app.business_owner or "", len(r.validation),
             round(r.effort_hours, 1), (r.reasons[0]["text"] if r.reasons else "")] for r in p.apps],
           color_cols=(3, 4))
    _sheet(wb, "Compatibility matrix",
           ["Area", "Okta feature", "PingFederate capability", "Compatibility", "Strategy", "Apps", "Tenant objects",
            "Evidence", "Review", "Sources"],
           [16, 34, 40, 13, 12, 34, 30, 60, 9, 50],
           [[m["cap"].area, m["cap"].okta_feature, m["cap"].pf_capability, m["cap"].compatibility, m["cap"].strategy,
             ", ".join(m["apps"]), ", ".join(m["objects"]), m["cap"].evidence, m["cap"].review,
             "\n".join(m["cap"].sources)] for m in p.matrix], color_cols=(4, 5))
    _sheet(wb, "Tenant objects",
           ["Kind", "Name", "Type", "Status", "Strategy", "Compatibility", "Used by apps", "Evidence", "Effort (h)"],
           [20, 32, 26, 10, 12, 13, 34, 60, 10],
           [[t["obj"].kind.replace("_", " "), t["obj"].name, t["obj"].subtype, t["obj"].status, t["obj"].strategy,
             t["obj"].compat_level, ", ".join(t["apps"]), "; ".join(h.evidence for h in t["hits"]), t["hours"]]
            for t in p.tenant], color_cols=(5, 6))
    _sheet(wb, "Dependencies",
           ["Depends", "On", "Kind", "Origin", "Status", "Evidence", "Decided by"],
           [34, 34, 16, 12, 12, 70, 18],
           [[d["source"], d["target"], d["dep"].kind, d["dep"].origin, d["dep"].status, d["dep"].evidence or "",
             d["dep"].decided_by or ""] for d in p.deps + p.suggested])
    wrows = []
    for w, rows in p.waves.items():
        for r in rows:
            wrows.append([w, r.app.label, r.app.protocol, r.strategy, r.risk, _d(p.timeline.cutover.get(w)),
                          "; ".join(r.wave_reasons + r.flags)])
    _sheet(wb, "Wave plan", ["Wave", "Application", "Protocol", "Strategy", "Risk", "Cutover (planned)", "Notes"],
           [22, 34, 9, 12, 10, 16, 60], wrows, color_cols=(4,))
    erows = []
    for r in p.apps:
        a = r.app
        nameid = ""
        if a.is_saml and a.saml:
            nameid = f"{a.saml.name_id_template or '?'} -> {a.saml.pf_nameid_source or '?'} {a.saml.pf_nameid_detail or ''}"
        attrs = "; ".join(f"{c.name}: {c.pf_source or '?'} {c.pf_source_detail or ''}".strip() for c in a.claims)
        erows.append([a.label, a.protocol, r.strategy, nameid, attrs, "\n".join(r.validation)])
    _sheet(wb, "Engineer detail", ["Application", "Protocol", "Strategy", "NameID (Okta -> PingFederate)",
                                   "Attributes (source in PingFederate)", "Validation checklist"],
           [30, 9, 12, 40, 60, 70], erows, color_cols=(3,))
    _sheet(wb, "Capability catalog",
           ["Key", "Area", "Okta feature", "PingFederate capability", "Compatibility", "Strategy", "Effort (h)",
            "Prerequisite", "Review", "Evidence", "Sources"],
           [26, 16, 34, 40, 13, 12, 9, 11, 9, 60, 50],
           [[c.key, c.area, c.okta_feature, c.pf_capability, c.compatibility, c.strategy, c.effort_hours,
             "yes" if c.prerequisite else "", c.review, c.evidence, "\n".join(c.sources)] for c in p.catalog.values()],
           color_cols=(5, 6))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_docx(p: StrategyPack, customer: str) -> bytes:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    s = p.summary
    d.add_paragraph(f"Migration Strategy Pack: Okta to {p.target}", style="Title")
    d.add_paragraph(f"{customer} · {p.generated:%d %b %Y} · catalog {p.catalog_version} · risk rules {p.rules_version}")

    d.add_heading("Executive summary", level=1)
    d.add_paragraph(
        f"Okta holds {s['total']} applications, {s['federated']} of them SAML or OIDC. {s['migratable']} can move to "
        f"{p.target} by recreating or transforming their configuration, {s['redesign']} need a redesign, "
        f"{s['retire']} are candidates to retire and {s['retain']} stay on Okta until a blocking decision is made. "
        f"{s['high_risk']} applications are high risk. The applications are planned in {s['waves']} wave(s), "
        f"finishing around {_d(s['end']) or 'a date to be agreed'}, after {s['foundation_hours']} hours of foundation "
        f"work on sign-in policies, MFA and other tenant-level designs.")
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Strategy", "Apps", "Meaning"
    for k in ("RECREATE", "TRANSFORM", "REDESIGN", "RETIRE", "RETAIN"):
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = k.title(), str(s[k.lower()]), STRATEGY_HELP[k]

    d.add_heading("Wave plan", level=1)
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Wave", "Applications", "Cutover"
    for w, rows in p.waves.items():
        c = t.add_row().cells
        c[0].text, c[1].text = w, ", ".join(r.app.label for r in rows)
        c[2].text = _d(p.timeline.cutover.get(w)) or "—"
    moved = [r for r in p.apps if r.wave_reasons or r.flags]
    if moved:
        d.add_paragraph("Sequencing notes (dependencies and strategy):")
        for r in moved:
            for x in r.wave_reasons + r.flags:
                d.add_paragraph(f"{r.app.label}: {x}", style="List Bullet")

    if p.foundation:
        d.add_heading("Foundation work before Wave 0", level=1)
        for f in p.foundation:
            d.add_paragraph(f"{f['obj'].name} ({f['obj'].kind.replace('_', ' ')}"
                            + (f", {f['obj'].subtype.replace('_', ' ').lower()}" if f['obj'].subtype else "")
                            + f"): {f['cap'].pf_capability}"
                            + (f" · used by {', '.join(f['apps'])}" if f["apps"] else ""), style="List Bullet")

    red = [r for r in p.apps if r.strategy == "REDESIGN"]
    if red:
        d.add_heading("Applications needing redesign", level=1)
        for r in red:
            d.add_paragraph(f"{r.app.label}: " + "; ".join(x["text"] for x in r.reasons[:3]), style="List Bullet")
    ret = [r for r in p.apps if r.strategy == "RETIRE"]
    if ret:
        d.add_heading("Retirement candidates (business validation needed)", level=1)
        for r in ret:
            d.add_paragraph(f"{r.app.label}: {r.reasons[0]['text'] if r.reasons else ''}", style="List Bullet")

    d.add_heading("Compatibility highlights", level=1)
    for m in [m for m in p.matrix if m["cap"].compatibility != "EQUIVALENT"][:15]:
        who = ", ".join((m["apps"] + m["objects"])[:4])
        d.add_paragraph(f"{m['cap'].okta_feature} → {m['cap'].pf_capability} ({m['cap'].compatibility.lower()}, "
                        f"{m['cap'].strategy.lower()}): {who}", style="List Bullet")

    if p.deps:
        d.add_heading("Key dependencies", level=1)
        for x in [x for x in p.deps if x["dep"].status == "CONFIRMED"][:25]:
            d.add_paragraph(f"{x['source']} depends on {x['target']} ({x['dep'].kind.replace('_', ' ').lower()})",
                            style="List Bullet")

    d.add_heading("Method", level=1)
    d.add_paragraph(
        "Compatibility, complexity, risk, dependencies and strategy come from deterministic, versioned rules "
        f"(capability catalog {p.catalog_version}, risk rules {p.rules_version}). AI agents review and explain these "
        "results but do not set them. Dependencies suggested from documents are used only after a named person "
        f"confirms them. {s['expert_rows']} catalog row(s) in use are marked for review by a PingFederate architect "
        "for this customer (licensing and version). Dates are planning estimates.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
