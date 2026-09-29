"""Pipeline: knowledge base, review agents, plan approval, gated PingFederate write,
SAML validation (captured + automated), cutover and rollback."""
import base64
import json
import threading
from pathlib import Path

import pytest

from app.config import Settings
from app.devtools.samlkit import build_response, make_keypair
from app.integrations.okta.client import build_source
from app.integrations.pingfederate.client import PfError
from app.integrations.pingfederate.writer import PfWriter, ref_paths
from app.models.db import AgentRun, Application, MigrationEvent, SessionLocal, init_engine
from app.models.state import MigrationState as S
from app.services import approval, cutover, decisions, knowledge, pf_write, risk
from app.services import saml_validation as sv
from app.services.agents import service as ag
from app.services.agents.schema import AgentReviewOutput
from app.services.discovery import run_discovery
from tests.conftest import csrf

KB_FILE = Path(__file__).resolve().parent.parent / "data" / "knowledge" / "northwind-sso-standards.md"
WIKI = "0oa3engwiki00000003"


@pytest.fixture
def disc(db, settings):
    run_discovery(db, build_source(settings), settings, actor="t")
    db.commit()
    return db


def _unblock(db, settings):
    decisions.decide(db, "OGNL_POLICY", "Lead", "Allow OGNL (reviewed, documented expressions)", None, None, None)
    decisions.decide(db, "DUPLICATE_ENTITY_ID", "Lead", "Decommission the duplicate (e.g. unused UAT)", None, None, None)
    decisions.decide(db, "OKTA_USER_ID_NAMEID", "Lead", "Copy Okta user IDs into a directory attribute", None, None, None)
    risk.score_all(db, risk.load_weights(settings.risk_weights_file))
    db.commit()


def _to_mapping_ready(db, settings, app):
    run = ag.run_panel(db, app, settings, "Asha")
    res = {k: "Checked with the vendor, fine" for k, _ in ag.disagreements(run)}
    ag.review(db, run, "Asha", "ACCEPTED", "", res)
    db.commit()
    return run


def _to_approved(db, settings, app):
    _to_mapping_ready(db, settings, app)
    p = approval.generate(db, app, settings, "Asha")
    approval.submit(db, p, "Asha")
    approval.decide(db, p, settings, "Ben", "APPROVED", "ok")
    db.commit()
    return p


# --- knowledge base ---------------------------------------------------------------
def test_kb_add_search_duplicate_and_formats(disc):
    doc = knowledge.add_document(disc, KB_FILE.name, KB_FILE.read_bytes(), "Asha")
    assert len(doc.chunks) >= 8
    hits = knowledge.search(disc, "SHA-1 signing algorithm")
    assert hits and hits[0].heading == "Signing" and hits[0].ref.startswith("KB-")
    with pytest.raises(knowledge.KnowledgeError, match="Already uploaded"):
        knowledge.add_document(disc, "copy.md", KB_FILE.read_bytes(), "Asha")
    with pytest.raises(knowledge.KnowledgeError, match="Unsupported"):
        knowledge.add_document(disc, "x.exe", b"MZ", "Asha")
    with pytest.raises(knowledge.KnowledgeError, match="name"):
        knowledge.add_document(disc, "y.txt", b"hello", "")
    html = b"<html><script>evil()</script><h2>Groups</h2><p>memberOf search for groups</p></html>"
    d2 = knowledge.add_document(disc, "g.html", html, "Asha")
    assert "evil" not in d2.chunks[0].text and d2.chunks[0].heading == "Groups"


def test_kb_docx_and_app_scoping(disc, tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Wiki vendor", 1)
    d.add_paragraph("Confluence needs the groups attribute named confluence-groups with prefix filter.")
    p = tmp_path / "wiki.docx"
    d.save(p)
    knowledge.add_document(disc, "wiki.docx", p.read_bytes(), "Asha", app_id=WIKI)
    assert knowledge.search(disc, "confluence-groups prefix filter", app_id=WIKI)
    assert not knowledge.search(disc, "confluence-groups prefix filter")   # app-specific doc hidden elsewhere


# --- agents -----------------------------------------------------------------------
def test_agents_offline_panel_valid_with_citations(disc, settings):
    knowledge.add_document(disc, KB_FILE.name, KB_FILE.read_bytes(), "Asha")
    app = disc.get(Application, WIKI)
    run = ag.run_panel(disc, app, settings, "Asha")
    assert run.status == "VALID" and [r.agent for r in run.reviews] == ["SAML_ANALYSIS", "CLAIMS_MAPPING", "GROUP_MAPPING", "RISK"]
    cited = {c for r in run.reviews for i in r.output["items"] for c in i["citations"]}
    assert cited and cited <= {ref for r in run.reviews for ref in r.knowledge_refs}
    assert ag.run_panel(disc, app, settings, "Asha").id == run.id          # cached on same input


def test_agent_policy_rejects_bad_output(disc, settings):
    app = disc.get(Application, WIKI)
    items = ag.build_items(disc, app, settings)["SAML_ANALYSIS"]
    ctx = ag.build_context(disc, app, "SAML_ANALYSIS", items, settings)
    good = [{"key": i.key, "assessment": "AGREE", "comment": "ok", "citations": []} for i in items]
    out = AgentReviewOutput.model_validate({"verdict": "AGREE", "summary": "fine", "items": good})
    assert ag.policy_errors(out, ctx) == []
    bad = good[1:] + [{"key": "made_up", "assessment": "DISAGREE", "comment": "see KB-999", "citations": ["KB-999"]}]
    out = AgentReviewOutput.model_validate({"verdict": "AGREE", "summary": "This is low risk", "items": bad})
    errs = " | ".join(ag.policy_errors(out, ctx))
    assert "does not exist: made_up" in errs and "Item not reviewed" in errs and "KB-999" in errs
    assert "Verdict AGREE does not match" in errs


class _FakeReviewer:
    name = "anthropic"
    model = "fake"

    def review(self, agent, context):
        from app.services.ai.providers import ProviderResult
        assert "_kb_by_item" not in context                      # internal hints never sent to the model
        return ProviderResult(raw={"verdict": "AGREE", "summary": "x", "items": [], "extra": 1}, model="fake")


def test_agent_schema_rejection_marks_run_rejected(disc, settings):
    run = ag.run_panel(disc, disc.get(Application, WIKI), settings, "Asha", reviewer=_FakeReviewer())
    assert run.status == "REJECTED" and all(r.status == "REJECTED_SCHEMA" for r in run.reviews)
    with pytest.raises(ValueError, match="all four"):
        ag.review(disc, run, "Asha", "ACCEPTED")


def test_agent_accept_needs_resolutions_and_moves_state(disc, settings):
    app = disc.get(Application, "0oa1expenseportal01")             # duplicate entity ID -> a DISAGREE
    run = ag.run_panel(disc, app, settings, "Asha")
    keys = [k for k, _ in ag.disagreements(run)]
    assert keys
    with pytest.raises(ValueError, match="resolved"):
        ag.review(disc, run, "Asha", "ACCEPTED")
    with pytest.raises(ValueError, match="named"):
        ag.review(disc, run, "claude", "ACCEPTED")
    ag.review(disc, run, "Asha", "ACCEPTED", "", {k: "UAT will be decommissioned" for k in keys})
    assert app.state == S.MAPPING_READY.value


# --- plan approval ------------------------------------------------------------------
def test_plan_generate_blocked_then_approve_four_eyes(disc, settings):
    app = disc.get(Application, WIKI)
    _to_mapping_ready(disc, settings, app)
    if app.blocked:
        with pytest.raises(ValueError, match="blocked"):
            approval.generate(disc, app, settings, "Asha")
        _unblock(disc, settings)
    p = approval.generate(disc, app, settings, "Asha")
    assert app.state == S.PLAN_GENERATED.value and p.plan["pingfederate"]["payload"]["active"] is False
    assert p.plan["rollback"]["idp_entity_id"] and p.plan["validation"]["audience"] == app.saml.audience
    approval.submit(disc, p, "Asha")
    assert app.state == S.AWAITING_APPROVAL.value
    with pytest.raises(ValueError, match="Four-eyes"):
        approval.decide(disc, p, settings, "asha", "APPROVED")
    with pytest.raises(ValueError, match="comment"):
        approval.decide(disc, p, settings, "Ben", "REJECTED")
    approval.decide(disc, p, settings, "Ben", "APPROVED", "Looks right")
    assert app.state == S.APPROVED.value and p.status == "APPROVED"


def test_plan_approval_refused_after_okta_drift(disc, settings):
    _unblock(disc, settings)
    app = disc.get(Application, WIKI)
    _to_mapping_ready(disc, settings, app)
    p = approval.generate(disc, app, settings, "Asha")
    approval.submit(disc, p, "Asha")
    app.saml.sso_acs_url = "https://wiki.northwind.example/new-acs"
    app.saml.acs_endpoints = []
    with pytest.raises(ValueError, match="changed since"):
        approval.decide(disc, p, settings, "Ben", "APPROVED")


# --- PingFederate write ---------------------------------------------------------------
class _Resp:
    def __init__(self, status, body):
        self.status_code, self._b = status, body
        self.text = json.dumps(body)

    def json(self):
        return self._b


class _FakeHttp:
    def __init__(self, existing=None, missing_refs=(), fail_post=False, timeout_post=False):
        self.headers, self.auth, self.posts = {}, None, []
        self.existing, self.missing, self.fail_post = existing or [], set(missing_refs), fail_post
        self.timeout_post = timeout_post

    def get(self, url, params=None, timeout=None, verify=None):
        if url.endswith("/idp/spConnections"):
            return _Resp(200, {"items": self.existing})
        return _Resp(404 if any(url.endswith(m) for m in self.missing) else 200, {})

    def post(self, url, json=None, timeout=None, verify=None):
        self.posts.append(json)
        if self.timeout_post:
            import requests
            raise requests.Timeout("read timed out")
        if self.fail_post:
            return _Resp(422, {"message": "Validation error(s) occurred.", "validationErrors": [
                {"fieldPath": "entityId", "message": "bad"}]})
        return _Resp(201, {**json, "id": "wiki-conn"})


@pytest.fixture
def write_settings(tmp_path, export_dir):
    return Settings(DATABASE_URL=f"sqlite:///{tmp_path / 'w.db'}", OKTA_SOURCE="file", OKTA_EXPORT_DIR=str(export_dir),
                    PF_SOURCE="live", PF_ADMIN_URL="https://pf.example:9999/pf-admin-api/v1", PF_ADMIN_USER="auditor",
                    PF_ADMIN_PASSWORD="r", PF_WRITE_ENABLED=True, PF_WRITE_USER="writer", PF_WRITE_PASSWORD="w",
                    PF_VARIABLES_FILE=str(tmp_path / "vars.json"), FLASK_SECRET_KEY="t", _env_file=None)


def _vars_for(plan, path: Path):
    names = pf_write.unresolved(plan.plan["pingfederate"]["payload"])
    path.write_text(json.dumps({n: {"description": "", "value": f"val-{n.lower()}"} for n in names}))


def test_write_requires_separate_account_and_live_source(tmp_path):
    with pytest.raises(ValueError, match="PF_WRITE_USER"):
        Settings(PF_SOURCE="live", PF_ADMIN_URL="https://x", PF_ADMIN_USER="a", PF_ADMIN_PASSWORD="b",
                 PF_WRITE_ENABLED=True, _env_file=None)
    with pytest.raises(ValueError, match="PF_SOURCE=live"):
        Settings(PF_WRITE_ENABLED=True, PF_WRITE_USER="w", PF_WRITE_PASSWORD="p", _env_file=None)


def test_writer_has_no_update_or_delete_and_refuses_active():
    assert not any(hasattr(PfWriter, m) for m in ("update_sp_connection", "delete_sp_connection", "put", "delete"))
    s = Settings(PF_SOURCE="live", PF_ADMIN_URL="https://x", PF_ADMIN_USER="a", PF_ADMIN_PASSWORD="b",
                 PF_WRITE_ENABLED=True, PF_WRITE_USER="w", PF_WRITE_PASSWORD="p", _env_file=None)
    w = PfWriter(s, session=_FakeHttp())
    assert w.http.auth == ("w", "p")                                 # write account, not the auditor
    with pytest.raises(PfError, match="disabled"):
        w.create_sp_connection({"active": True})


def test_dry_run_apply_creates_disabled_connection(write_settings, tmp_path):
    init_engine(write_settings.database_url)
    db = SessionLocal()
    run_discovery(db, build_source(write_settings), write_settings, actor="t")
    _unblock(db, write_settings)
    app = db.get(Application, WIKI)
    plan = _to_approved(db, write_settings, app)
    http = _FakeHttp()
    w = PfWriter(write_settings, session=http)
    d = pf_write.dry_run(db, app, write_settings, "Asha", writer=w)
    assert not d.ok and d.unresolved                                  # variables not filled yet
    with pytest.raises(ValueError, match="Checks failed"):
        pf_write.apply(db, app, write_settings, "Asha", writer=w)
    _vars_for(plan, Path(write_settings.pf_variables_file))
    with pytest.raises(ValueError, match="dry run first"):
        pf_write.apply(db, app, write_settings, "Asha", writer=w)
    d = pf_write.dry_run(db, app, write_settings, "Asha", writer=w)
    assert d.ok, d.checks
    assert set(ref_paths(d.payload)) == {"IdP adapter", "LDAP data store", "Signing key pair"}
    obj = pf_write.apply(db, app, write_settings, "Asha", writer=w)
    assert len(http.posts) == 1 and http.posts[0]["active"] is False and "{{" not in json.dumps(http.posts[0])
    assert obj.pf_id == "wiki-conn" and app.state == S.PF_CONFIGURED.value
    SessionLocal.remove()


def test_dry_run_blocks_existing_entity_and_missing_refs(write_settings):
    init_engine(write_settings.database_url)
    db = SessionLocal()
    run_discovery(db, build_source(write_settings), write_settings, actor="t")
    _unblock(db, write_settings)
    app = db.get(Application, WIKI)
    plan = _to_approved(db, write_settings, app)
    _vars_for(plan, Path(write_settings.pf_variables_file))
    http = _FakeHttp(existing=[{"entityId": app.saml.audience, "name": "Wiki (manual)"}],
                     missing_refs=["/keyPairs/signing/val-signing_key_pair_id"])
    d = pf_write.dry_run(db, app, write_settings, "Asha", writer=PfWriter(write_settings, session=http))
    failed = {label for label, ok, _ in d.checks if not ok}
    assert failed == {"Entity ID not yet on PingFederate", "Signing key pair exists"}
    SessionLocal.remove()


def test_write_failure_is_audited_and_state_unchanged(write_settings):
    init_engine(write_settings.database_url)
    db = SessionLocal()
    run_discovery(db, build_source(write_settings), write_settings, actor="t")
    _unblock(db, write_settings)
    app = db.get(Application, WIKI)
    plan = _to_approved(db, write_settings, app)
    _vars_for(plan, Path(write_settings.pf_variables_file))
    w = PfWriter(write_settings, session=_FakeHttp(fail_post=True))
    pf_write.dry_run(db, app, write_settings, "Asha", writer=w)
    with pytest.raises(PfError, match="entityId: bad"):
        pf_write.apply(db, app, write_settings, "Asha", writer=w)
    db.rollback()                                                    # what the route does on an error
    assert db.get(Application, WIKI).state == S.APPROVED.value
    ev = {e.event_type: e for e in db.query(MigrationEvent).all()}
    assert "PF_WRITE_ATTEMPT" in ev and not ev["PF_WRITE_FAILED"].detail["may_have_been_created"]
    # A timeout: the POST may have landed, so the error says so and the audit keeps it.
    w = PfWriter(write_settings, session=_FakeHttp(timeout_post=True))
    pf_write.dry_run(db, app, write_settings, "Asha", writer=w)
    with pytest.raises(PfError, match="may or may not"):
        pf_write.apply(db, app, write_settings, "Asha", writer=w)
    db.rollback()
    last = db.query(MigrationEvent).filter_by(event_type="PF_WRITE_FAILED").order_by(MigrationEvent.id.desc()).first()
    assert last.detail["may_have_been_created"] is True
    SessionLocal.remove()


def test_manual_import_needs_object_on_pingfederate(disc, settings):
    _unblock(disc, settings)
    app = disc.get(Application, WIKI)
    _to_approved(disc, settings, app)
    from app.models.db import PfObject
    for o in disc.query(PfObject).all():                            # make sure nothing matches yet
        disc.delete(o)
    disc.flush()
    with pytest.raises(ValueError, match="No matching connection"):
        pf_write.confirm_manual(disc, app, "Asha")
    disc.add(PfObject(kind="SP_CONNECTION", pf_id="wiki", key=app.saml.audience, name="Wiki", active=False,
                      details={"acs_urls": [], "attributes": []}))
    disc.flush()
    pf_write.confirm_manual(disc, app, "Asha")
    assert app.state == S.PF_CONFIGURED.value


# --- SAML validation -------------------------------------------------------------------
@pytest.fixture(scope="module")
def keys():
    return make_keypair("pf"), make_keypair("okta"), make_keypair("attacker")


def _resp_for(app, settings, key, cert, issuer=None, name_id="jane@northwind.example", attrs=None, **kw):
    exp = cutover.expected_values(app, settings)
    values = attrs if attrs is not None else {n: ([f"v-{n}"] if n not in exp["group_attributes"] else ["Engineering"])
                                              for n in exp["attributes"]}
    return build_response(issuer or exp["issuer"], exp["audience"], exp["acs_urls"][0], name_id, values,
                          name_id_format=exp["name_id_format"] or "urn:oasis:names:tc:SAML:1.1:nameid-format:unspecified",
                          key_pem=key, cert_pem=cert, **kw)


def test_validate_pass_fail_and_tamper(disc, settings, keys):
    (pk, pc), _, (ak, ac) = keys
    app = disc.get(Application, WIKI)
    exp = cutover.expected_values(app, settings)
    good = _resp_for(app, settings, pk, pc)
    rep = sv.validate(good, exp, signing_cert_pem=pc)
    assert rep.verdict == "PASS", [(c.key, c.status, c.note) for c in rep.checks if c.status != "PASS"]
    assert sv.validate(good, exp, signing_cert_pem=ac).verdict == "FAIL"             # wrong certificate
    forged = _resp_for(app, settings, ak, ac)                                         # attacker key, own KeyInfo cert
    r = sv.validate(forged, exp, signing_cert_pem=pc)
    assert r.verdict == "FAIL" and any(c.key == "cert:assertion" for c in r.checks)
    xml = base64.b64decode(good).replace(b"jane@northwind.example</saml:NameID>", b"eve@northwind.example</saml:NameID>")
    assert sv.validate(base64.b64encode(xml).decode(), exp, signing_cert_pem=pc).verdict == "FAIL"
    wrong_aud = build_response(exp["issuer"], "https://other", exp["acs_urls"][0], "x", {}, key_pem=pk, cert_pem=pc)
    r = sv.validate(wrong_aud, exp, signing_cert_pem=pc)
    assert {c.key for c in r.checks if c.status == "FAIL"} >= {"audience"}
    assert sv.validate(good, exp).verdict == "PASS_WITH_WARNINGS"                     # no cert -> not verified


def test_validate_input_forms_and_xxe(disc, settings, keys):
    (pk, pc), _, _ = keys
    app = disc.get(Application, WIKI)
    exp = cutover.expected_values(app, settings)
    b64 = _resp_for(app, settings, pk, pc)
    from urllib.parse import quote_plus
    assert sv.validate(f"SAMLResponse={quote_plus(b64)}&RelayState=x", exp, signing_cert_pem=pc).verdict == "PASS"
    assert sv.validate(base64.b64decode(b64).decode(), exp, signing_cert_pem=pc).verdict == "PASS"
    with pytest.raises(sv.SamlInputError, match="DOCTYPE"):
        sv.decode('<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><samlp:Response/>')
    with pytest.raises(sv.SamlInputError):
        sv.validate("not saml at all", exp)
    err = build_response(exp["issuer"], exp["audience"], exp["acs_urls"][0], "x", {}, sign="none",
                         status="urn:oasis:names:tc:SAML:2.0:status:Responder")
    assert sv.validate(err, exp).verdict == "FAIL"


def _cfg(settings, cert_pem, tmp_path):
    """Settings with the PingFederate signing certificate configured (the only one that counts for the gate)."""
    f = tmp_path / "pf-signing.crt"
    f.write_text(cert_pem)
    return settings.model_copy(update={"pf_signing_cert_file": f})


def test_baseline_comparison(disc, settings, keys, tmp_path):
    (pk, pc), (ok_, oc), _ = keys
    settings = _cfg(settings, pc, tmp_path)
    app = disc.get(Application, WIKI)
    exp = cutover.expected_values(app, settings)
    names = exp["attributes"]
    okta = _resp_for(app, settings, ok_, oc, issuer=cutover.okta_issuer(app),
                     attrs={n: ["Engineering"] if n in exp["group_attributes"] else ["Jane"] for n in names})
    base = cutover.record_validation(disc, app, settings, okta, "Asha", "BASELINE_OKTA", cert_pem=oc)
    assert base.verdict in ("PASS", "PASS_WITH_WARNINGS"), base.result["checks"]
    pf_resp = _resp_for(app, settings, pk, pc,
                        attrs={n: ["Engineering"] if n in exp["group_attributes"] else ["JANE"] for n in names})
    v = cutover.record_validation(disc, app, settings, pf_resp, "Asha", "PRE_CUTOVER")
    assert v.result["baseline_used"]
    statuses = {c["key"]: c["status"] for c in v.result["checks"]}
    non_group = [n for n in names if n not in exp["group_attributes"]]
    if non_group:
        assert statuses[f"baseline:{non_group[0]}"] == "WARN"                 # case-only difference
    assert v.response_sha256 and pf_resp[:60] not in json.dumps(v.result)     # raw response not stored


# --- cutover / rollback ---------------------------------------------------------------
def _pf_configured(db, settings, app):
    _unblock(db, settings)
    _to_approved(db, settings, app)
    from app.models.db import PfObject
    db.add(PfObject(kind="SP_CONNECTION", pf_id="wiki-x", key=app.saml.audience, name="Wiki", active=True,
                    details={"acs_urls": [{"url": u} for u in cutover.expected_values(app, settings)["acs_urls"]],
                             "attributes": cutover.expected_values(app, settings)["attributes"] + ["SAML_SUBJECT"],
                             "sign_assertions": True, "has_issuance_criteria": True,
                             "virtual_identities": [app.saml.idp_issuer],
                             "slo_urls": [app.saml.slo_logout_url] if app.saml.slo_enabled else []}))
    db.flush()
    pf_write.confirm_manual(db, app, "Asha")


def test_full_cutover_then_post_validation(disc, settings, keys, tmp_path):
    (pk, pc), _, _ = keys
    settings = _cfg(settings, pc, tmp_path)
    app = disc.get(Application, WIKI)
    _pf_configured(disc, settings, app)
    with pytest.raises(ValueError, match="passing pre-cutover"):
        cutover.business_signoff(disc, app, settings, "Owner", "")
    v = cutover.record_validation(disc, app, settings, _resp_for(app, settings, pk, pc), "Asha", "PRE_CUTOVER")
    assert v.verdict == "PASS" and app.state == S.TESTING.value
    cutover.business_signoff(disc, app, settings, "Owner", "Roles fine")
    gate = cutover.gate(disc, app, settings)
    assert cutover.gate_ok(gate), [(c.label, c.detail) for c in gate if c.required and not c.ok]
    cutover.go_decision(disc, app, settings, "Lead", "CHG-1")
    with pytest.raises(ValueError, match="Record"):
        cutover.cutover_done(disc, app, "Lead", "", settings)
    cutover.cutover_done(disc, app, "Lead", "CHG-1 switched 10:02", settings)
    assert app.state == S.MIGRATED.value
    cutover.record_validation(disc, app, settings, _resp_for(app, settings, pk, pc), "Asha", "POST_CUTOVER")
    assert app.state == S.VALIDATED.value
    rb = cutover.runbook(disc, app, settings)
    assert rb["rollback_values"]["idp_entity_id"] and cutover.runbook_docx(rb, "Northwind")[:2] == b"PK"


def test_failed_post_validation_then_rollback(disc, settings, keys, tmp_path):
    (pk, pc), _, (ak, ac) = keys
    settings = _cfg(settings, pc, tmp_path)
    app = disc.get(Application, WIKI)
    _pf_configured(disc, settings, app)
    cutover.record_validation(disc, app, settings, _resp_for(app, settings, pk, pc), "Asha", "PRE_CUTOVER")
    cutover.business_signoff(disc, app, settings, "Owner", "ok")
    cutover.go_decision(disc, app, settings, "Lead", "")
    cutover.cutover_done(disc, app, "Lead", "switched")
    v = cutover.record_validation(disc, app, settings, _resp_for(app, settings, ak, ac), "Asha", "POST_CUTOVER")
    assert v.verdict == "FAIL" and app.state == S.FAILED.value
    with pytest.raises(Exception, match="named"):
        cutover.start_rollback(disc, app, settings, "system", "x")
    cutover.start_rollback(disc, app, settings, "Lead", "Signature invalid at SP")
    assert app.state == S.ROLLBACK.value
    ev = disc.query(MigrationEvent).filter_by(app_id=app.id, to_state="ROLLBACK").one()
    assert ev.detail["restore"]["idp_entity_id"] and "certificate_pem" not in ev.detail["restore"]
    cutover.confirm_rollback(disc, app, "Lead", "Test user signed in via Okta")
    assert app.state == S.OKTA_ACTIVE.value


def test_gate_red_without_validation(disc, settings):
    app = disc.get(Application, WIKI)
    _pf_configured(disc, settings, app)
    gate = cutover.gate(disc, app, settings)
    assert not cutover.gate_ok(gate)
    assert {c.label for c in gate if c.required and not c.ok} >= {"Pre-cutover validation passed (current plan, since any failure)", "Business sign-off"}


# --- automated test SSO (headless browser against a local mock IdP) ------------------------
def _mock_idp(response_b64: str, acs: str):
    from flask import Flask, request
    idp = Flask("mock-idp")

    @idp.get("/idp/startSSO.ping")
    def start():
        return ('<form method="post" action="/login"><input name="pf.username"><input name="pf.pass" type="password">'
                '<button id="signOnButton" type="submit">Sign on</button></form>')

    @idp.post("/login")
    def login():
        assert request.form["pf.username"] == "test.user"
        return (f'<form id="f" method="post" action="{acs}"><input type="hidden" name="SAMLResponse" value="{response_b64}">'
                '</form><script>document.getElementById("f").submit()</script>')
    from werkzeug.serving import make_server
    srv = make_server("127.0.0.1", 0, idp)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def test_automated_test_sso_captures_and_validates(disc, settings, keys, tmp_path):
    pytest.importorskip("playwright.sync_api")
    (pk, pc), _, _ = keys
    app = disc.get(Application, WIKI)
    resp = _resp_for(app, settings, pk, pc)
    srv = _mock_idp(resp, cutover.expected_values(app, settings)["acs_urls"][0])
    try:
        s = settings.model_copy(update={
            "test_sso_enabled": True, "test_sso_username": "test.user",
            "test_sso_password": __import__("pydantic").SecretStr("pw"),
            "test_sso_start_url": f"http://127.0.0.1:{srv.server_port}/idp/startSSO.ping?PartnerSpId={{entity_id}}",
            "test_sso_timeout_seconds": 20, "pf_signing_cert_file": _cfg(settings, pc, tmp_path).pf_signing_cert_file})
        from app.services import test_sso
        cap = test_sso.capture(s, app.saml.audience, screenshot_dir=tmp_path)
        assert "SAMLResponse=" in cap.form_body
        v = cutover.record_validation(disc, app, s, cap.form_body, "Asha", "PRE_CUTOVER", "AUTOMATED")
        assert v.verdict == "PASS" and v.source == "AUTOMATED"
    finally:
        srv.shutdown()


def test_test_sso_off_by_default(settings):
    from app.services import test_sso
    with pytest.raises(test_sso.TestSsoError, match="off"):
        test_sso.capture(settings, "x")


# --- routes ---------------------------------------------------------------------------------
def test_pipeline_pages_and_flow_via_ui(client, app, settings, keys):
    (pk, pc), _, _ = keys
    tok = csrf(client)
    client.post("/discovery/run", data={"_csrf": tok, "actor": "Asha"})
    assert client.get("/pipeline").status_code == 200
    assert client.get("/knowledge").status_code == 200
    r = client.post("/knowledge/upload", data={"_csrf": tok, "actor": "Asha",
                                               "file": (KB_FILE.open("rb"), KB_FILE.name)},
                    content_type="multipart/form-data", follow_redirects=True)
    assert b"passages" in r.data
    assert b"KB-" in client.get("/knowledge?q=SHA-1").data
    db = SessionLocal()
    _unblock(db, settings)
    SessionLocal.remove()
    page = client.get(f"/applications/{WIKI}/pipeline")
    assert page.status_code == 200 and b"Run the 4 agents" in page.data
    client.post(f"/applications/{WIKI}/agents/run", data={"_csrf": tok, "actor": "Asha", "provider": "offline"})
    db = SessionLocal()
    run = db.query(AgentRun).filter_by(app_id=WIKI).one()
    form = {"_csrf": tok, "actor": "Asha", "decision": "ACCEPTED"}
    form.update({f"resolve::{k}": "ok" for k, _ in ag.disagreements(run)})
    SessionLocal.remove()
    client.post(f"/agent-runs/{run.id}/review", data=form)
    client.post(f"/applications/{WIKI}/plan/generate", data={"_csrf": tok, "actor": "Asha"})
    db = SessionLocal()
    plan = approval.latest_plan(db, WIKI)
    SessionLocal.remove()
    client.post(f"/plans/{plan.id}/submit", data={"_csrf": tok, "actor": "Asha"})
    r = client.post(f"/plans/{plan.id}/decide", data={"_csrf": tok, "actor": "Asha", "decision": "APPROVED"},
                    follow_redirects=True)
    assert b"Four-eyes" in r.data
    client.post(f"/plans/{plan.id}/decide", data={"_csrf": tok, "actor": "Ben", "decision": "APPROVED"})
    assert client.get(f"/plans/{plan.id}.json").json["status"] == "APPROVED"
    r = client.post(f"/applications/{WIKI}/pf/dry-run", data={"_csrf": tok, "actor": "Asha"}, follow_redirects=True)
    assert b"Writes enabled" in r.data                                   # dry run shows writes are off
    db = SessionLocal()
    a = db.get(Application, WIKI)
    assert a.state == "APPROVED"
    resp = _resp_for(a, settings, pk, pc)
    SessionLocal.remove()
    r = client.post(f"/applications/{WIKI}/validate", data={"_csrf": tok, "actor": "Asha", "stage": "PRE_CUTOVER",
                                                           "response": resp, "cert_pem": pc}, follow_redirects=True)
    assert b"Validation recorded" in r.data
    db = SessionLocal()
    vid = cutover.results(db, WIKI)[0].id
    SessionLocal.remove()
    assert client.get(f"/validations/{vid}").status_code == 200
    assert client.get(f"/applications/{WIKI}/runbook.docx").data[:2] == b"PK"
    assert client.get(f"/applications/0oa7employeeportal7/pipeline").status_code == 404   # OIDC: not in this pipeline


# --- safety review regressions ------------------------------------------------------------
def test_signature_must_be_verified_with_configured_cert_for_the_gate(disc, settings, keys, tmp_path):
    (pk, pc), _, (ak, ac) = keys
    app = disc.get(Application, WIKI)
    _pf_configured(disc, settings, app)
    # No certificate configured: an unverifiable signature fails in the gating stages.
    v = cutover.record_validation(disc, app, settings, _resp_for(app, settings, ak, ac), "Asha", "PRE_CUTOVER")
    assert v.verdict == "FAIL" and {c["key"]: c["status"] for c in v.result["checks"]}["signature_valid"] == "FAIL"
    # An attacker's response verified with a pasted (attacker's) certificate never counts for the gate.
    v = cutover.record_validation(disc, app, settings, _resp_for(app, settings, ak, ac), "Asha", "PRE_CUTOVER",
                                  cert_pem=ac)
    assert v.result["cert_source"] == "pasted" and not cutover.counts_for_gate(v)
    with pytest.raises(ValueError, match="configured PingFederate"):
        cutover.business_signoff(disc, app, settings, "Owner", "")
    # A configured certificate that is missing on disk is an error, not a silent skip.
    s = settings.model_copy(update={"pf_signing_cert_file": tmp_path / "nope.crt"})
    with pytest.raises(ValueError, match="not found"):
        cutover.record_validation(disc, app, s, _resp_for(app, settings, pk, pc), "Asha", "PRE_CUTOVER")


def test_replay_old_response_and_retest_after_failure(disc, settings, keys, tmp_path):
    from datetime import timedelta
    from app.models.db import utcnow
    (pk, pc), _, (ak, ac) = keys
    settings = _cfg(settings, pc, tmp_path)
    app = disc.get(Application, WIKI)
    _pf_configured(disc, settings, app)
    pre = _resp_for(app, settings, pk, pc)
    cutover.record_validation(disc, app, settings, pre, "Asha", "PRE_CUTOVER")
    cutover.business_signoff(disc, app, settings, "Owner", "ok")
    cutover.go_decision(disc, app, settings, "Lead", "")
    cutover.cutover_done(disc, app, "Lead", "switched", settings)
    with pytest.raises(ValueError, match="already recorded"):               # same response re-used as proof
        cutover.record_validation(disc, app, settings, pre, "Asha", "POST_CUTOVER")
    old = _resp_for(app, settings, pk, pc, now=utcnow() - timedelta(hours=2))
    with pytest.raises(ValueError, match="before the cutover"):             # captured before the switch
        cutover.record_validation(disc, app, settings, old, "Asha", "POST_CUTOVER")
    cutover.record_validation(disc, app, settings, _resp_for(app, settings, ak, ac), "Asha", "POST_CUTOVER")
    assert app.state == S.FAILED.value
    cutover.retry_testing(disc, app, "Lead", "fixed the certificate")
    with pytest.raises(ValueError, match="pre-cutover"):                    # the old pass no longer counts
        cutover.business_signoff(disc, app, settings, "Owner", "")
    assert not cutover.gate_ok(cutover.gate(disc, app, settings))
    cutover.record_validation(disc, app, settings, _resp_for(app, settings, pk, pc), "Asha", "PRE_CUTOVER")
    cutover.business_signoff(disc, app, settings, "Owner", "re-tested")


def test_rollback_possible_after_validated_within_window(disc, settings, keys, tmp_path):
    (pk, pc), _, _ = keys
    settings = _cfg(settings, pc, tmp_path)
    app = disc.get(Application, WIKI)
    _pf_configured(disc, settings, app)
    cutover.record_validation(disc, app, settings, _resp_for(app, settings, pk, pc), "Asha", "PRE_CUTOVER")
    cutover.business_signoff(disc, app, settings, "Owner", "ok")
    cutover.go_decision(disc, app, settings, "Lead", "")
    cutover.cutover_done(disc, app, "Lead", "switched", settings)
    cutover.record_validation(disc, app, settings, _resp_for(app, settings, pk, pc), "Asha", "POST_CUTOVER")
    assert app.state == S.VALIDATED.value
    cutover.declare_failed(disc, app, "Lead", "Payroll role missing on day 2")
    cutover.start_rollback(disc, app, settings, "Lead", "roles")
    assert app.state == S.ROLLBACK.value


def test_encrypted_assertion_fails_in_strict_mode():
    xml = (b'<samlp:Response xmlns:samlp="urn:oasis:names:tc:SAML:2.0:protocol" '
           b'xmlns:saml="urn:oasis:names:tc:SAML:2.0:assertion" ID="r1" IssueInstant="2026-01-01T00:00:00Z">'
           b'<saml:Issuer>https://pf</saml:Issuer><samlp:Status><samlp:StatusCode '
           b'Value="urn:oasis:names:tc:SAML:2.0:status:Success"/></samlp:Status>'
           b'<saml:EncryptedAssertion><x/></saml:EncryptedAssertion></samlp:Response>')
    exp = {"issuer": "https://pf", "acs_urls": [], "attributes": []}
    assert {c.key: c.status for c in sv.validate(xml, exp).checks}["encrypted"] == "WARN"
    assert {c.key: c.status for c in sv.validate(xml, exp, strict=True).checks}["encrypted"] == "FAIL"
