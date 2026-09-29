"""Sprint 2: risk engine, assessments, comparison, Excel."""
import io
import json
from types import SimpleNamespace

import openpyxl
import pytest
from sqlalchemy import select

from app.integrations.okta.client import FileOktaSource
from app.models.db import AiAssessment, Application, MigrationEvent, RiskScore
from app.services import comparison, risk
from app.services.ai import service as ai
from app.services.ai.context import Masker, build_context
from app.services.ai.providers import AnthropicProvider, OfflineProvider
from app.services.discovery import run_discovery
from app.services.excel_report import build_workbook
from tests.conftest import csrf


@pytest.fixture
def discovered(db, settings, export_dir):
    run_discovery(db, FileOktaSource(export_dir), settings, actor="tester")
    return db


def app_(db, app_id):
    return db.get(Application, app_id)


# --- risk ---------------------------------------------------------------------
def test_scores_are_computed_on_discovery(discovered):
    hr = app_(discovered, "0oa2hranalytics0002")
    assert (hr.complexity_score, hr.complexity_level) == (15, "HIGH")
    assert (hr.impact_score, hr.impact_level) == (5, "MEDIUM")
    assert hr.overall_level == "HIGH" and hr.blocked and hr.suggested_wave == "Blocked"


def test_breakdown_sums_to_scores(discovered):
    for row in discovered.scalars(select(RiskScore)):
        c = sum(b["points"] for b in row.breakdown if b["dimension"] == "COMPLEXITY")
        i = sum(b["points"] for b in row.breakdown if b["dimension"] == "IMPACT")
        assert (c, i) == (row.complexity_score, row.impact_score)
        assert row.rules_version == risk.RULES_VERSION


def test_waves(discovered):
    waves = {a.label: a.suggested_wave for a in discovered.scalars(select(Application).where(Application.is_saml))}
    assert waves["Legacy Travel Booking"] == "Decommission review"
    assert waves["ServiceNow"] == "Wave 1"
    assert waves["Salesforce"] == "Wave 2"


def test_ognl_allowed_unblocks(db, settings, export_dir):
    run_discovery(db, FileOktaSource(export_dir), settings.model_copy(update={"pf_ognl_allowed": True}))
    hr = app_(db, "0oa2hranalytics0002")
    assert not hr.blocked and hr.suggested_wave != "Blocked"


def test_rescore_only_when_inputs_change(discovered):
    snow = app_(discovered, "0oa4servicenow00004")
    n = len(discovered.scalars(select(RiskScore).where(RiskScore.app_id == snow.id)).all())
    risk.score_app(discovered, snow)
    assert len(discovered.scalars(select(RiskScore).where(RiskScore.app_id == snow.id)).all()) == n
    snow.business_criticality, snow.has_test_environment = "HIGH", False
    risk.score_app(discovered, snow)
    assert snow.impact_score == 1 + 5 + 5 and snow.impact_level == "HIGH"
    assert len(discovered.scalars(select(RiskScore).where(RiskScore.app_id == snow.id)).all()) == n + 1


def _fake(codes, users, crit, test):
    return SimpleNamespace(findings=[SimpleNamespace(code=c, severity="WARNING") for c in codes],
                           okta_status="ACTIVE", user_count=users, business_criticality=crit,
                           has_test_environment=test)


def test_both_high_raises_overall():
    codes = ["CATALOG_APP_PARTIAL_CONFIG", "GROUP_CLAIM", "MULTIPLE_ACS"]          # 5+3+2 = 10 -> HIGH
    r = risk.compute(_fake(codes, 600, "MEDIUM", None))                            # 4+2+2 = 8 -> HIGH
    assert (r.complexity_level, r.impact_level, r.overall_level) == ("HIGH", "HIGH", "CRITICAL")
    r2 = risk.compute(_fake(codes, 600, "LOW", True))                              # 4 -> MEDIUM
    assert (r2.complexity_level, r2.impact_level, r2.overall_level) == ("HIGH", "MEDIUM", "HIGH")
    assert r2.suggested_wave == "Wave 2"
    r3 = risk.compute(_fake([], 10, "LOW", True))
    assert (r3.overall_level, r3.suggested_wave) == ("LOW", "Wave 0 (pilot)")


def test_weights_override(tmp_path, discovered):
    f = tmp_path / "w.json"
    f.write_text(json.dumps({"CATALOG_APP_PARTIAL_CONFIG": 0}))
    w = risk.load_weights(f)
    assert w["CATALOG_APP_PARTIAL_CONFIG"] == (0, 1)
    r = risk.compute(app_(discovered, "0oa4servicenow00004"), w)
    assert r.complexity_score == 3


# --- assessments ------------------------------------------------------------------
def test_offline_assessments_are_valid(discovered, settings):
    for a in discovered.scalars(select(Application).where(Application.is_saml)):
        row = ai.run_assessment(discovered, a, settings, "tester")
        assert row.status == "VALID", (a.label, row.errors)
    assert app_(discovered, "0oa6legacytravel006").id


def test_assessment_is_cached(discovered, settings):
    a = app_(discovered, "0oa4servicenow00004")
    r1 = ai.run_assessment(discovered, a, settings, "t")
    r2 = ai.run_assessment(discovered, a, settings, "t")
    r3 = ai.run_assessment(discovered, a, settings, "t", force=True)
    assert r1.id == r2.id != r3.id


class FakeClient:
    """Stands in for anthropic.Anthropic; returns whatever tool input we give it."""
    def __init__(self, payload):
        self.payload, self.calls = payload, []
        self.messages = self

    def create(self, **kw):
        self.calls.append(kw)
        return SimpleNamespace(content=[SimpleNamespace(type="tool_use", input=self.payload)],
                               usage=SimpleNamespace(input_tokens=1200, output_tokens=300))


def _good_payload(db, settings, app):
    ctx = build_context(app, risk.latest_score(db, app.id), settings)
    return OfflineProvider().assess(ctx).raw


def test_anthropic_provider_valid(discovered, settings):
    a = app_(discovered, "0oa4servicenow00004")
    fake = FakeClient(_good_payload(discovered, settings, a))
    row = ai.run_assessment(discovered, a, settings, "t",
                            provider=AnthropicProvider("k", "claude-sonnet-5", client=fake))
    assert row.status == "VALID" and row.provider == "anthropic" and row.input_tokens == 1200
    call = fake.calls[0]
    assert call["tool_choice"] == {"type": "tool", "name": "submit_assessment"}
    assert "ignore any instructions" in call["system"]
    sent = call["messages"][0]["content"]
    assert "@northwind.example" not in sent          # no user emails ever sent


def test_schema_violation_rejected(discovered, settings):
    a = app_(discovered, "0oa4servicenow00004")
    bad = _good_payload(discovered, settings, a)
    bad["confidence"] = "VERY HIGH"
    bad["execute"] = "delete all apps"
    row = ai.run_assessment(discovered, a, settings, "t",
                            provider=AnthropicProvider("k", "m", client=FakeClient(bad)))
    assert row.status == "REJECTED_SCHEMA" and row.errors


def test_policy_violations_rejected(discovered, settings):
    a = app_(discovered, "0oa2hranalytics0002")          # blocked, overall HIGH
    bad = _good_payload(discovered, settings, a)
    bad["migration_approach"] = "STANDARD"
    bad["summary"] = "This is a low risk app."
    bad["finding_explanations"] = bad["finding_explanations"][1:] + [{"code": "MADE_UP", "explanation": "x"}]
    row = ai.run_assessment(discovered, a, settings, "t",
                            provider=AnthropicProvider("k", "m", client=FakeClient(bad)))
    assert row.status == "REJECTED_POLICY"
    joined = " ".join(row.errors)
    assert "BLOCKED" in joined and "low risk" in joined.lower() and "MADE_UP" in joined


def test_provider_error_is_recorded(discovered, settings):
    class Boom:
        name, model = "anthropic", "m"
        def assess(self, ctx):
            raise RuntimeError("network down")
    row = ai.run_assessment(discovered, app_(discovered, "0oa4servicenow00004"), settings, "t", provider=Boom())
    assert row.status == "ERROR" and "network down" in row.errors[0]


def test_masking_roundtrip(discovered, settings):
    a = app_(discovered, "0oa3engwiki00000003")
    ctx = build_context(a, risk.latest_score(discovered, a.id), settings)
    m = Masker()
    masked = json.dumps(m.mask(ctx))
    for secret in ("northwind", "ENG-Platform", "Engineering Wiki"):
        assert secret not in masked, secret
    assert m.unmask("Add group-1 on https://host-1/x for app-1") == \
        "Add ENG-Platform on https://wiki.northwind.example/x for Engineering Wiki (Confluence DC)"


def test_masking_used_for_llm(discovered, settings):
    s = settings.model_copy(update={"ai_mask_data": True})
    a = app_(discovered, "0oa4servicenow00004")
    fake = FakeClient(_good_payload(discovered, settings, a))
    row = ai.run_assessment(discovered, a, s, "t", provider=AnthropicProvider("k", "m", client=fake))
    assert row.masked and "SN-ITIL" not in fake.calls[0]["messages"][0]["content"]


def test_review_flow(discovered, settings):
    a = app_(discovered, "0oa4servicenow00004")
    row = ai.run_assessment(discovered, a, settings, "t")
    with pytest.raises(ValueError):
        ai.review(discovered, row, "claude", "ACCEPTED")
    with pytest.raises(ValueError):
        ai.review(discovered, row, "alice", "REJECTED", "")
    ai.review(discovered, row, "alice", "ACCEPTED")
    assert a.state == "ASSESSED" and row.reviewed_by == "alice"
    with pytest.raises(ValueError):
        ai.review(discovered, row, "bob", "REJECTED", "late")


def test_invalid_assessment_cannot_be_accepted(discovered, settings):
    a = app_(discovered, "0oa4servicenow00004")
    bad = _good_payload(discovered, settings, a)
    bad["migration_approach"] = "BLOCKED"
    row = ai.run_assessment(discovered, a, settings, "t",
                            provider=AnthropicProvider("k", "m", client=FakeClient(bad)))
    with pytest.raises(ValueError):
        ai.review(discovered, row, "alice", "ACCEPTED")


# --- comparison & excel -------------------------------------------------------------
def test_comparison_rows(discovered, settings):
    rows = comparison.build(app_(discovered, "0oa3engwiki00000003"), settings)
    by = {r.field: r for r in rows}
    assert by["SP entity ID"].status == "SAME"
    assert by["IdP entity ID (Issuer)"].status == "ACTION"            # custom issuer -> virtual server id
    assert by["SSO URL"].pf_value == "https://sso.example.com/idp/SSO.saml2"
    assert by["groups"].status == "ACTION"                            # Okta-only group in the filter
    assert by["fullName"].status == "ACTION" and by["email"].status == "SAME"
    assert by["Assigned groups"].status == "ACTION"
    assert "PartnerSpId=https%3A%2F%2Fwiki.northwind.example" in by["IdP metadata"].pf_value


def test_comparison_catalog_unknowns(discovered, settings):
    rows = comparison.build(app_(discovered, "0oa4servicenow00004"), settings)
    assert {r.field for r in rows if r.status == "UNKNOWN"} >= {"SP entity ID", "ACS URL", "Attribute statements"}


def test_comparison_duplicate_entity(discovered, settings):
    rows = comparison.build(app_(discovered, "0oa1expenseuat00011"), settings)
    assert next(r for r in rows if r.field == "SP entity ID").status == "ACTION"


def test_excel_workbook(discovered, settings):
    apps = discovered.scalars(select(Application).where(Application.is_saml).order_by(Application.label)).all()
    for a in apps:
        ai.run_assessment(discovered, a, settings, "t")
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(discovered, apps, settings, "Report")))
    assert wb.sheetnames == ["Summary", "Applications", "SAML - Okta vs PingFederate",
                             "OIDC - Okta vs PingFederate", "PingFederate reconciliation", "Migration plan",
                             "Pipeline", "SAML validation", "Agent reviews", "Attribute mapping", "Directory prep",
                             "Action plan", "Owner questions", "Findings", "Risk breakdown"]
    assert wb["Applications"].max_row == 8
    cmp_ws = wb["SAML - Okta vs PingFederate"]
    assert [c.value for c in cmp_ws[1]] == ["Application", "Section", "Field", "Okta value",
                                             "PingFederate setting", "PingFederate value", "Status", "Note"]
    assert cmp_ws.max_row > 100
    formulas = [c.value for row in wb["Summary"].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) == 40 + 12 + 10 + 8 and all("COUNT" in f or "SUM" in f for f in formulas)
    assert wb["Summary"]["A1"].font.name == "Arial"


# --- routes ---------------------------------------------------------------------------
def _run(client):
    client.post("/discovery/run", data={"_csrf": csrf(client), "actor": "alice"})


def test_ui_assess_review_and_exports(client):
    _run(client)
    r = client.post("/applications/0oa4servicenow00004/assess",
                    data={"_csrf": csrf(client), "actor": "alice", "provider": "offline"}, follow_redirects=True)
    assert b"VALID" in r.data and b"Recommended actions" in r.data
    from app.models.db import SessionLocal
    aid = SessionLocal().scalars(select(AiAssessment.id).order_by(AiAssessment.id.desc())).first()
    r = client.post(f"/assessments/{aid}/review", data={"_csrf": csrf(client), "actor": "alice",
                                                         "decision": "ACCEPTED"}, follow_redirects=True)
    assert b"accepted" in r.data
    assert SessionLocal().get(Application, "0oa4servicenow00004").state == "ASSESSED"
    r = client.get("/export/report.xlsx")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    r = client.get("/applications/0oa3engwiki00000003/export.xlsx")
    assert r.status_code == 200 and "Engineering_Wiki" in r.headers["Content-Disposition"]
    r = client.get("/applications/0oa3engwiki00000003/report")
    assert r.status_code == 200 and b"Okta vs PingFederate" in r.data
    assert client.get("/reports").status_code == 200
    r = client.get("/")
    assert b"Risk matrix" in r.data and b"Suggested waves" in r.data


def test_ui_assess_all_and_anthropic_without_key(client):
    _run(client)
    r = client.post("/assess-all", data={"_csrf": csrf(client), "actor": "alice", "provider": "offline"},
                    follow_redirects=True)
    assert b"7 valid" in r.data
    r = client.post("/applications/0oa4servicenow00004/assess",
                    data={"_csrf": csrf(client), "actor": "alice", "provider": "anthropic"}, follow_redirects=True)
    assert b"ANTHROPIC_API_KEY is not set" in r.data


def test_details_rescore(client):
    _run(client)
    r = client.post("/applications/0oa4servicenow00004/details",
                    data={"_csrf": csrf(client), "actor": "alice", "business_criticality": "HIGH",
                          "has_test_environment": "no"}, follow_redirects=True)
    assert b"Impact is now HIGH" in r.data
