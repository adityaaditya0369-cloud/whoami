"""OIDC discovery/comparison, protocol tabs and the side-by-side comparison view."""
import io

import openpyxl
import pytest
from sqlalchemy import select

from app.integrations.okta.client import FileOktaSource
from app.models.db import Application, Finding
from app.routes.web import _diffmark
from app.services import comparison
from app.services.discovery import run_discovery
from app.services.excel_report import build_workbook
from tests.conftest import csrf


@pytest.fixture
def discovered(db, settings, export_dir):
    run_discovery(db, FileOktaSource(export_dir), settings, actor="tester")
    return db


def codes(db, app_id):
    return {f.code for f in db.scalars(select(Finding).where(Finding.app_id == app_id))}


def test_oidc_parsed(discovered):
    a = discovered.get(Application, "0oa7employeeportal7")
    assert a.is_oidc and not a.is_saml and a.protocol == "OIDC"
    assert a.oidc.client_id == "0oa7EmpPortalClient01"
    assert a.oidc.grant_types == ["authorization_code", "refresh_token"]
    assert a.user_count == 25 and a.group_count == 1


@pytest.mark.parametrize("app_id,expected,absent", [
    ("0oa7employeeportal7", {"OIDC_NEW_CLIENT_SECRET", "OIDC_REFRESH_TOKENS", "OIDC_ACCESS_CONTROL_REQUIRED",
                             "OIDC_ISSUER_CHANGE"}, {"ISSUANCE_CRITERIA_REQUIRED", "OIDC_IMPLICIT_GRANT"}),
    ("0oa8mobileexpense08", {"OIDC_REFRESH_TOKENS", "OKTA_NATIVE_GROUPS_ASSIGNED"},
     {"OIDC_PKCE_NOT_REQUIRED", "OIDC_NEW_CLIENT_SECRET"}),
    ("0oa9reportingapi009", {"OIDC_MACHINE_CLIENT", "OIDC_PRIVATE_KEY_JWT"}, {"NO_ASSIGNMENTS"}),
    ("0oa10legacyspa00010", {"OIDC_IMPLICIT_GRANT", "OIDC_WILDCARD_REDIRECT", "DIRECT_USER_ASSIGNMENTS"}, set()),
])
def test_oidc_findings(discovered, app_id, expected, absent):
    found = codes(discovered, app_id)
    assert expected <= found, expected - found
    assert not (absent & found)


def test_oidc_scored(discovered):
    rpt = discovered.get(Application, "0oa9reportingapi009")
    assert rpt.overall_level and rpt.suggested_wave != "Decommission review"
    spa = discovered.get(Application, "0oa10legacyspa00010")
    assert spa.complexity_level == "HIGH"


def test_oidc_comparison(discovered, settings):
    rows = comparison.build(discovered.get(Application, "0oa7employeeportal7"), settings)
    by = {r.field: r for r in rows}
    assert by["Client ID"].status == "SAME" and by["Client ID"].pf_value == "0oa7EmpPortalClient01"
    assert by["Client authentication"].status == "ACTION"
    assert by["Token endpoint"].pf_value == "https://sso.example.com/as/token.oauth2"
    assert by["Existing refresh tokens"].status == "SP_UPDATE"
    assert by["Scopes"].status == "UNKNOWN"
    assert comparison.sections(rows) == ["OAuth client", "Redirects", "Grants & tokens", "Endpoints",
                                         "Scopes & claims", "Access control"]
    spa = {r.field: r for r in comparison.build(discovered.get(Application, "0oa10legacyspa00010"), settings)}
    assert spa["Grant: implicit"].status == "ACTION" and spa["Redirect URI 1"].status == "ACTION"
    svc = {r.field for r in comparison.build(discovered.get(Application, "0oa9reportingapi009"), settings)}
    assert "Authorization endpoint" not in svc and "Token endpoint" in svc


def test_excel_has_oidc_sheet(discovered, settings):
    apps = discovered.scalars(select(Application).where(
        Application.is_saml | Application.is_oidc).order_by(Application.label)).all()
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(discovered, apps, settings, "R")))
    o = wb["OIDC - Okta vs PingFederate"]
    assert o.max_row > 30 and {c.value for c in o["A"][1:]} == {
        "Employee Portal", "Mobile Expenses", "Reporting API", "Legacy Intranet SPA"}
    protos = [c.value for c in wb["Applications"]["C"][1:]]
    assert protos.count("SAML") == 7 and protos.count("OIDC") == 4


def test_diffmark_highlights_changes():
    out = str(_diffmark("https://sso.example.com/idp/SSO.saml2", "https://northwind.okta.com/app/x/sso/saml", "new"))
    assert '<mark class="new">' in out
    assert str(_diffmark("same", "same")) == "same"
    assert "&lt;script&gt;" in str(_diffmark("<script>x</script>", "abc", "new"))   # escaped


def test_protocol_tabs(client):
    client.post("/discovery/run", data={"_csrf": csrf(client), "actor": "a"})
    r = client.get("/applications?protocol=OIDC")
    assert b"Employee Portal" in r.data and b"Engineering Wiki" not in r.data and b"Grant types" in r.data
    r = client.get("/applications?protocol=OTHER")
    assert b"Company Handbook" in r.data and b"Employee Portal" not in r.data
    r = client.get("/applications")
    assert b"Engineering Wiki" in r.data and b"Employee Portal" not in r.data
    assert client.get("/applications?scope=all").status_code == 200
    r = client.get("/?protocol=OIDC")
    assert r.status_code == 200 and b"PingFederate readiness (SAML)" not in r.data
    assert b"PingFederate readiness (SAML)" in client.get("/?protocol=SAML").data


def test_side_by_side_view(client):
    client.post("/discovery/run", data={"_csrf": csrf(client), "actor": "a"})
    r = client.get("/applications/0oa3engwiki00000003")
    assert b'class="diff"' in r.data and b"Show only differences" in r.data and b"<mark" in r.data
    r = client.get("/applications/0oa7employeeportal7")
    assert r.status_code == 200 and b"OAuth client" in r.data and b"Token endpoint" in r.data and b"oauth2" in r.data
    assert b"OIDC client" in r.data and b'data-tab="claims"' not in r.data
    assert client.get("/applications/0oa7employeeportal7/report").status_code == 200
    r = client.get("/export/report.xlsx?protocol=OIDC")
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb["Applications"].max_row == 5
