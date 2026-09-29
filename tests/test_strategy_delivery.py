"""Usage analysis, decisions/what-if, estimator, strategy report, comms pack, build package."""
import io
import json
import shutil
import time
import zipfile
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document
from sqlalchemy import select

from app.config import Settings
from app.integrations.okta.client import FileOktaSource, LiveOktaClient, UsageUnavailable
from app.models.db import Application, Decision, MigrationEvent
from app.services import build_package as bpk
from app.services import comms
from app.services import decisions as dm
from app.services import estimate as em
from app.services import risk
from app.services import strategy as stm
from app.services.discovery import run_discovery
from tests.conftest import csrf


@pytest.fixture
def disc(db, settings, export_dir):
    run_discovery(db, FileOktaSource(export_dir), settings, actor="t")
    return db


def app_(db, i):
    return db.get(Application, i)


def codes(a):
    return {f.code for f in a.findings}


# --- 1. usage -------------------------------------------------------------------------
def test_usage_collected(disc):
    exp = app_(disc, "0oa1expenseportal01")
    assert exp.usage_known and exp.usage_unique_users == 22 and exp.usage_events == 132 and exp.usage_last_seen
    assert "UNUSED_IN_WINDOW" in codes(app_(disc, "0oa1expenseuat00011"))
    assert "UNUSED_IN_WINDOW" in codes(app_(disc, "0oa10legacyspa00010"))
    assert "UNUSED_IN_WINDOW" not in codes(app_(disc, "0oa9reportingapi009"))       # machine client
    assert app_(disc, "0oa1expenseuat00011").suggested_wave == "Decommission review"
    score = risk.latest_score(disc, "0oa1expenseportal01")
    assert any(b["rule"] == "ACTIVE_USERS" for b in score.breakdown)


def test_usage_missing_is_not_fatal(db, settings, export_dir):
    for p in export_dir.glob("apps/*/logs.json"):
        p.unlink()
    run = run_discovery(db, FileOktaSource(export_dir), settings)
    assert run.status == "SUCCESS" and not app_(db, "0oa1expenseportal01").usage_known
    assert not any("UNUSED_IN_WINDOW" in codes(a) for a in db.scalars(select(Application)))


class FakeResp:
    def __init__(self, status=200, body=None, text=""):
        self.status_code, self._b, self.text, self.headers, self.links = status, body, text, {}, {}

    def json(self):
        return self._b


class FakeSession:
    def __init__(self, rs):
        self.rs, self.calls, self.headers = list(rs), [], {}

    def get(self, url, params=None, headers=None, timeout=None):
        self.calls.append(params)
        return self.rs.pop(0)


def _ssws():
    return Settings(OKTA_SOURCE="live", OKTA_ORG_URL="https://x.okta.com", OKTA_AUTH_MODE="ssws",
                    OKTA_API_TOKEN="t", _env_file=None)


def test_system_log_stops_on_short_page_and_filters():
    fs = FakeSession([FakeResp(body=[{"uuid": i} for i in range(1000)]), FakeResp(body=[{"uuid": "x"}])])
    fs.rs[0].links = {"next": {"url": "https://x.okta.com/api/v1/logs?after=1"}}
    ev = LiveOktaClient(_ssws(), fs).list_app_usage_events("0oaA", 90, 5000)
    assert len(ev) == 1001 and 'target.id eq "0oaA"' in fs.calls[0]["filter"]


def test_system_log_forbidden_is_usage_unavailable(monkeypatch):
    monkeypatch.setattr(time, "sleep", lambda s: None)
    fs = FakeSession([FakeResp(403, text="forbidden")])
    with pytest.raises(UsageUnavailable):
        LiveOktaClient(_ssws(), fs).list_app_usage_events("0oaA", 90, 5000)


# --- 2. decisions & what-if ---------------------------------------------------------------
def test_register_and_decide_unblocks(disc):
    reg = {v.d.key: v for v in dm.register(disc)}
    assert {a.label for a in reg["OGNL_POLICY"].blocked_apps} == {"Engineering Wiki (Confluence DC)", "HR Analytics (Tableau)"}
    assert "DUPLICATE_ENTITY_ID" in reg and "DECOMMISSION_UNUSED" in reg
    with pytest.raises(ValueError):
        dm.decide(disc, "OGNL_POLICY", "claude", reg["OGNL_POLICY"].d.options[0], None, None, None)
    with pytest.raises(ValueError):
        dm.decide(disc, "OGNL_POLICY", "alice", "Maybe", None, None, None)
    dm.decide(disc, "OGNL_POLICY", "alice", reg["OGNL_POLICY"].d.options[1], None, "2026-10-10", "SteerCo #3")
    risk.score_all(disc)
    hr = app_(disc, "0oa2hranalytics0002")
    assert not hr.blocked and hr.suggested_wave != "Blocked"
    assert disc.scalars(select(MigrationEvent).where(MigrationEvent.event_type == "DECISION_RECORDED")).first()
    dm.decide(disc, "OGNL_POLICY", "alice", None, None, None, None, reopen=True)
    risk.score_all(disc)
    assert app_(disc, "0oa2hranalytics0002").blocked


def test_what_if_does_not_persist(disc):
    sim = dm.simulate(disc, ["OGNL_POLICY", "DUPLICATE_ENTITY_ID", "DECOMMISSION_UNUSED"])
    moved = {a.label: after for a, _b, after in sim.moved}
    assert moved["HR Analytics (Tableau)"] != "Blocked" and moved["Expense Portal (UAT)"] == "Out of scope"
    assert app_(disc, "0oa2hranalytics0002").blocked          # nothing saved
    assert disc.scalars(select(Decision).where(Decision.status == "DECIDED")).first() is None


# --- 4. estimator -------------------------------------------------------------------------------
def test_app_effort_is_deterministic(disc):
    e = em.app_effort(app_(disc, "0oa4servicenow00004"))
    assert e.build == 4.0 + 2 + 0.5 and e.test == 4.5 and "CATALOG_APP_PARTIAL_CONFIG" in e.drivers


def test_timeline_and_assumptions(disc):
    apps = dm.scope_apps(disc)
    t1 = em.build_timeline(disc, apps)
    assert set(t1.waves) == {"Wave 1", "Wave 2", "Blocked"}          # decommission candidates excluded
    assert all(p.end >= p.start for p in t1.phases)
    assert t1.cutover["Blocked"] > t1.cutover["Wave 1"]
    em.update_settings(disc, "alice", {"engineers": "0.5", "vendor_lead_days": "30", "start_date": "2026-10-05"})
    t2 = em.build_timeline(disc, apps)
    assert t2.end > t1.end and t2.start.strftime("%Y-%m-%d") == "2026-10-05"
    with pytest.raises(ValueError):
        em.update_settings(disc, "alice", {"engineers": "two"})
    with pytest.raises(ValueError):
        em.update_settings(disc, "", {})


# --- 3. strategy ----------------------------------------------------------------------------------
def test_strategy_numbers(disc, settings):
    st = stm.build(disc, settings)
    assert st.scope["okta_total"] == 13 and st.scope["in_scope"] == 11 and st.scope["decommission"] == 3
    assert st.scope["to_migrate"] == 8
    assert st.decisions_open[0].d.key in ("OGNL_POLICY", "DUPLICATE_ENTITY_ID")
    assert any("CATALOG" in r["risk"].upper() or "catalog" in r["risk"] for r in st.risks)
    assert str(st.scope["to_migrate"]) in st.narrative["executive_summary"]
    doc = Document(io.BytesIO(stm.to_docx(st, settings)))
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert heads[:3] == ["Executive summary", "Scope", "Approach"] and "Next steps (two weeks)" in heads


def test_cert_expiry_before_cutover_is_a_risk(disc, settings):
    em.update_settings(disc, "alice", {"start_date": "2026-11-01"})
    hr = app_(disc, "0oa2hranalytics0002")
    hr.wave = "Wave 1"                      # expires 2026-11-11; Wave 1 cutover is later than that
    st = stm.build(disc, settings)
    assert any("expires 2026-11-11" in r["risk"] for r in st.risks)


class FakeClient:
    def __init__(self, payload):
        self.payload, self.messages = payload, self

    def create(self, **kw):
        return SimpleNamespace(content=[SimpleNamespace(type="tool_use", input=self.payload)])


def test_ai_narrative_must_not_invent_numbers(disc, settings):
    st = stm.build(disc, settings)
    good = {"executive_summary": f"We will migrate {st.scope['to_migrate']} applications.", "key_messages": [],
            "recommendation": "Decide now."}
    row = stm.generate_draft(disc, st, settings, "alice", client=FakeClient(good))
    assert row.status == "VALID"
    assert stm.build(disc, settings).draft is not None
    bad = {**good, "executive_summary": "We will migrate 57 applications and save 43% of cost."}
    row = stm.generate_draft(disc, st, settings, "alice", client=FakeClient(bad))
    assert row.status == "REJECTED_POLICY" and "57" in row.errors[0]


# --- 5. comms ---------------------------------------------------------------------------------------
def test_comms_messages(disc, settings):
    tl = em.build_timeline(disc, dm.scope_apps(disc))
    wiki = comms.build_message(app_(disc, "0oa3engwiki00000003"), settings, tl, ["Nested groups?"])
    assert "https://sso.example.com/idp/SSO.saml2" in wiki.body and "Nested groups?" in wiki.body
    assert "Production cutover" in wiki.body and wiki.cutover
    snow = comms.build_message(app_(disc, "0oa4servicenow00004"), settings, tl, [])
    assert "SAML SP metadata" in snow.body
    emp = comms.build_message(app_(disc, "0oa7employeeportal7"), settings, tl, [], "Secrets vault hand-off")
    assert "Client ID stays the same: 0oa7EmpPortalClient01" in emp.body
    assert "never by email" in emp.body and "Secrets vault hand-off" in emp.body
    assert "/as/token.oauth2" in emp.body
    wb = openpyxl.load_workbook(io.BytesIO(comms.to_xlsx([wiki, snow, emp])))
    assert wb.active.max_row == 4 and wb.active["D1"].value == "Subject"


# --- 6. build package -----------------------------------------------------------------------------
def test_sp_connection_json(disc, settings):
    p = bpk.sp_connection(app_(disc, "0oa3engwiki00000003"), settings, ognl_allowed=False)
    j = p.payload
    assert j["entityId"] == "https://wiki.northwind.example" and j["active"] is False
    assert [e["url"] for e in j["spBrowserSso"]["ssoServiceEndpoints"]][1].startswith("https://wiki-dr")
    names = {a["name"] for a in j["spBrowserSso"]["attributeContract"]["extendedAttributes"]}
    assert names == {"email", "fullName", "wikiRole", "groups"}
    ful = j["spBrowserSso"]["adapterMappings"][0]["attributeContractFulfillment"]
    assert ful["email"]["value"] == "mail" and ful["fullName"]["value"].startswith("{{TODO_")
    assert "issuanceCriteria" in j["spBrowserSso"]["adapterMappings"][0]
    assert j["virtualIdentities"] == ["https://sso.northwind.example/wiki"]
    assert any("OGNL" in t or "Pre-compute" in t for t in p.todos)
    p2 = bpk.sp_connection(app_(disc, "0oa3engwiki00000003"), settings, ognl_allowed=True)
    assert p2.payload["spBrowserSso"]["adapterMappings"][0]["attributeContractFulfillment"]["fullName"]["source"]["type"] == "EXPRESSION"


def test_oauth_client_json(disc):
    p = bpk.oauth_client(app_(disc, "0oa10legacyspa00010"), modernise_implicit=True)
    assert p.payload["grantTypes"] == ["AUTHORIZATION_CODE"] and p.payload["requireProofKeyForCodeExchange"]
    p = bpk.oauth_client(app_(disc, "0oa9reportingapi009"), modernise_implicit=False)
    assert p.payload["clientAuth"]["type"] == "PRIVATE_KEY_JWT" and "jwksSettings" in p.payload
    assert "oidcPolicy" not in p.payload
    p = bpk.oauth_client(app_(disc, "0oa7employeeportal7"), False)
    assert "clientSecret" not in json.dumps(p.payload) and p.payload["clientId"] == "0oa7EmpPortalClient01"


def test_zip_package(disc, settings):
    apps = [a for a in dm.scope_apps(disc) if a.suggested_wave != "Decommission review"]
    data = bpk.to_zip(bpk.build_all(apps, settings, False, False), apps, settings, "All")
    z = zipfile.ZipFile(io.BytesIO(data))
    names = set(z.namelist())
    assert {"TODO.md", "README.md", "variables.json", "directory/create-groups.ps1"} <= names
    for n in names:
        if n.endswith(".json"):
            json.loads(z.read(n))
    ps1 = z.read("directory/create-groups.ps1").decode()
    assert "$w = -not $Apply" in ps1 and "-WhatIf:$w" in ps1 and "Remove-" not in ps1
    assert "PF-Engineering-Wiki-Confluence-DC-Users" in ps1
    assert "LDAP_DATA_STORE_ID" in json.loads(z.read("variables.json"))


# --- routes --------------------------------------------------------------------------------------------
def test_strategy_and_delivery_routes(client):
    client.post("/discovery/run", data={"_csrf": csrf(client), "actor": "alice"})
    r = client.get("/strategy")
    for t in (b"Executive summary", b"Decisions", b"What-if", b"Timeline", b"Decommission candidates"):
        assert t in r.data, t
    r = client.get("/strategy?assume=OGNL_POLICY&assume=DECOMMISSION_UNUSED")
    assert b"Apps that move" in r.data and b"Out of scope" in r.data
    r = client.post("/decisions/DUPLICATE_ENTITY_ID", data={"_csrf": csrf(client), "actor": "alice",
                    "choice": "Decommission the duplicate (e.g. unused UAT)"}, follow_redirects=True)
    assert b"Decision saved" in r.data
    r = client.post("/strategy/assumptions", data={"_csrf": csrf(client), "actor": "alice", "engineers": "3"},
                    follow_redirects=True)
    assert b"Assumptions saved" in r.data
    assert client.get("/strategy.docx").data[:2] == b"PK"
    r = client.get("/delivery")
    assert b"App-owner messages" in r.data and b"Build package" in r.data and b"create-groups.ps1" in r.data
    assert client.get("/delivery/comms.xlsx").data[:2] == b"PK"
    assert zipfile.ZipFile(io.BytesIO(client.get("/delivery/build-package.zip?wave=Wave 1").data)).namelist()
    assert client.get("/delivery/build-package.zip?app=0oa4servicenow00004").status_code == 200
    r = client.post("/strategy/narrative", data={"_csrf": csrf(client), "actor": "alice"}, follow_redirects=True)
    assert b"ANTHROPIC_API_KEY is not set" in r.data
